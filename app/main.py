"""Kochi Coastal One Health Observatory: API and dashboard server."""
import csv
import io
import math
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from . import analytics, config, idsp, reports, rules, sources, sync_odk, translate, weather
from .db import SessionLocal, get_db, init_db
from .models import (Advisory, AuditLog, Decision, HomeVisit, Household, HouseholdSurvey, IdspAlert, IdspCase,
                     LabResult, Person, PersonPII, Registration, Report, ReportForm, Site, StationRecord, SyncState,
                     User, WaterSample)
from .security import ROLES, Principal, authenticate, current_user, hash_password, make_token, require

ALL = ROLES
ANALYSTS = ("admin", "coordinator", "scientist", "doctor", "mo", "dso")
IDSP_VIEW = ("admin", "coordinator", "scientist", "doctor", "asha", "jphn", "mo", "dso", "lab")
IDSP_VERIFY = ("admin", "asha", "jphn", "mo")          # health staff confirm community signals
IDSP_ALERT = ("admin", "jphn", "mo", "dso")
IDSP_EXPORT = ("admin", "coordinator", "mo", "dso")
MANAGERS = ("admin", "coordinator")

@asynccontextmanager
async def lifespan(_app):
    init_db()
    db = SessionLocal()
    try:
        reports.ensure_defaults(db)
    finally:
        db.close()
    yield


app = FastAPI(title="Kochi Coastal One Health Observatory", version="1.0.0", lifespan=lifespan,
              docs_url="/api/docs" if config.ENVIRONMENT != "production" else None, redoc_url=None)


@app.middleware("http")
async def audit_and_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; "
        "script-src 'self'; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
    if config.ENVIRONMENT == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    if request.url.path.startswith("/api/") and request.url.path != "/api/health":
        db = SessionLocal()
        try:
            db.add(AuditLog(username=getattr(request.state, "username", None), method=request.method,
                            path=str(request.url.path)[:300], status=response.status_code,
                            ip=request.client.host if request.client else None))
            db.commit()
        finally:
            db.close()
    return response


def sites_for(db: Session, p: Principal, include_test: bool):
    allowed = None if p.sites is None else p.sites
    return [s.code for s in analytics.site_codes(db, include_test, allowed)]


# ------------------------------------------------------------------ auth

class Login(BaseModel):
    username: str
    password: str


@app.post("/api/auth/login")
def login(body: Login, db: Session = Depends(get_db)):
    user = authenticate(db, body.username.strip(), body.password)
    if not user:
        raise HTTPException(401, "Wrong username or password.")
    return {"token": make_token(user), "user": {"username": user.username, "name": user.full_name,
                                                 "role": user.role, "sites": user.sites}}


@app.get("/api/me")
def me(p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    u = db.get(User, p.username)
    return {"username": u.username, "name": u.full_name, "role": u.role, "sites": u.sites, "pii": p.pii}


@app.get("/api/health")
def health():
    return {"status": "ok", "time": datetime.now(timezone.utc).isoformat()}


# ------------------------------------------------------------------ overview

@app.get("/api/sites")
def sites(include_test: bool = False, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    codes = sites_for(db, p, include_test)
    fc = analytics.forecast(db)
    return [analytics.site_risk(db, db.get(Site, c), fc=fc) for c in codes]


@app.get("/api/summary")
def summary(include_test: bool = False, days: int = 7, p: Principal = Depends(current_user),
            db: Session = Depends(get_db)):
    codes = sites_for(db, p, include_test)
    since = date.today().toordinal() - days
    since = date.fromordinal(since)
    samples = db.query(WaterSample).filter(WaterSample.site_code.in_(codes), WaterSample.collected_on >= since).count()
    lab = (db.query(LabResult).join(WaterSample, WaterSample.sample_id == LabResult.sample_id)
           .filter(WaterSample.site_code.in_(codes), LabResult.method == "nio_lab", LabResult.tested_on >= since).count())
    return {
        "days": days,
        "streams": [
            {"source": "Water samples", "value": samples, "what": "collected"},
            {"source": "CSIR-NIO lab", "value": lab, "what": "results received"},
            {"source": "Camp registrations", "value": db.query(Registration).filter(
                Registration.site_code.in_(codes), Registration.collected_on >= since).count(), "what": "people"},
            {"source": "Station records", "value": db.query(StationRecord).filter(
                StationRecord.site_code.in_(codes), StationRecord.collected_on >= since).count(), "what": "screenings"},
            {"source": "Household visits", "value": db.query(HomeVisit).filter(
                HomeVisit.site_code.in_(codes), HomeVisit.visit_on >= since).count(), "what": "by ASHAs and volunteers"},
            {"source": "Household reports", "value": db.query(Report).filter(
                Report.site_code.in_(codes), Report.reported_on >= since).count(), "what": "illness and One Health"},
            {"source": "Advisories", "value": db.query(Advisory).filter(
                Advisory.site_code.in_(codes), Advisory.status == "sent").count(), "what": "sent to households"},
        ],
        "sync": [{"form": s.form_id, "last_run": s.last_run.isoformat() if s.last_run else None, "total": s.count}
                 for s in db.query(SyncState)],
    }


@app.get("/api/forecast")
def get_forecast(p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    return analytics.forecast(db)


class Scenario(BaseModel):
    rain_mm: float = Field(80, ge=0, le=500)
    tide_m: float = Field(1.1, ge=0, le=3)
    temp_c: float = Field(32, ge=15, le=50)
    rh: float = Field(78, ge=5, le=100)


@app.post("/api/model/scenario")
def model_scenario(body: Scenario, include_test: bool = False, p: Principal = Depends(current_user),
                   db: Session = Depends(get_db)):
    return analytics.scenario(db, body.rain_mm, body.tide_m, body.temp_c, body.rh, sites_for(db, p, include_test))


# ------------------------------------------------------------------ statistics

@app.get("/api/stats/trend")
def stats_trend(include_test: bool = False, weeks: int = Query(8, ge=2, le=104),
                p: Principal = Depends(require(*ANALYSTS)), db: Session = Depends(get_db)):
    return analytics.weekly_trend(db, sites_for(db, p, include_test), weeks)


@app.get("/api/stats/salinity-bp")
def stats_salinity(include_test: bool = False, p: Principal = Depends(require(*ANALYSTS)), db: Session = Depends(get_db)):
    return analytics.salinity_bp(db, sites_for(db, p, include_test))


@app.get("/api/stats/associations")
def stats_assoc(include_test: bool = False, p: Principal = Depends(require(*ANALYSTS)), db: Session = Depends(get_db)):
    return analytics.associations(db, sites_for(db, p, include_test))


@app.get("/api/stats/model")
def stats_model(include_test: bool = False, p: Principal = Depends(require(*ANALYSTS)), db: Session = Depends(get_db)):
    return analytics.sample_model(db, sites_for(db, p, include_test))


@app.get("/api/quality")
def get_quality(include_test: bool = False, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    return analytics.quality(db, sites_for(db, p, include_test))


@app.get("/api/loop")
def get_loop(include_test: bool = False, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    return analytics.loop_closed(db, sites_for(db, p, include_test))


@app.get("/api/ageing")
def get_ageing(include_test: bool = False, p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    return analytics.ageing(db, sites_for(db, p, include_test))


# ------------------------------------------------------------------ actions

@app.get("/api/tasks")
def tasks(include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "asha", "doctor")),
          db: Session = Depends(get_db)):
    return rules.asha_tasks(db, sites_for(db, p, include_test), p.pii)


@app.post("/api/rules/run")
def run_rules(include_test: bool = False, p: Principal = Depends(require(*MANAGERS)), db: Session = Depends(get_db)):
    return rules.generate(db, sites_for(db, p, include_test))


@app.get("/api/decisions")
def decisions(include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "ward")),
              db: Session = Depends(get_db)):
    codes = sites_for(db, p, include_test)
    rows = db.query(Decision).filter(Decision.site_code.in_(codes)).order_by(Decision.id.desc()).all()
    return [{"id": d.id, "site": d.site_code, "title": d.title, "reason": d.reason, "status": d.status,
             "decided_by": d.decided_by} for d in rows]


class StatusBody(BaseModel):
    status: str


@app.post("/api/decisions/{did}")
def decide(did: int, body: StatusBody, p: Principal = Depends(require("admin", "coordinator", "ward")),
           db: Session = Depends(get_db)):
    d = db.get(Decision, did)
    if not d or not p.can_see_site(d.site_code):
        raise HTTPException(404, "Not found.")
    if body.status not in ("open", "approved", "declined"):
        raise HTTPException(422, "Status must be open, approved or declined.")
    d.status, d.decided_by, d.decided_at = body.status, p.username, datetime.utcnow()
    db.commit()
    return {"ok": True}


@app.get("/api/advisories")
def advisories(include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "doctor", "asha")),
               db: Session = Depends(get_db)):
    codes = sites_for(db, p, include_test)
    rows = db.query(Advisory).filter(Advisory.site_code.in_(codes)).order_by(Advisory.id.desc()).all()
    return [{"id": a.id, "site": a.site_code, "hh_id": a.hh_id, "kind": a.kind, "text_en": a.text_en,
             "text_ml": a.text_ml, "status": a.status, "reviewed_by": a.reviewed_by} for a in rows]


class CustomAdvisory(BaseModel):
    site: str
    text_en: str = Field(min_length=5, max_length=1500)
    text_ml: str | None = Field(default=None, max_length=3000)
    hh_id: str | None = None


@app.post("/api/advisories/custom")
def custom_advisory(body: CustomAdvisory, p: Principal = Depends(require("admin", "coordinator", "doctor")),
                    db: Session = Depends(get_db)):
    """A message written by the team. Malayalam comes from Sarvam AI when configured, unless typed in."""
    if not db.get(Site, body.site) or not p.can_see_site(body.site):
        raise HTTPException(404, "Unknown site.")
    ml = (body.text_ml or "").strip() or translate.to_malayalam(db, body.text_en)
    key = f"custom:{body.site}:{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')}"
    db.add(Advisory(rule_key=key, site_code=body.site, hh_id=body.hh_id, kind="custom", text_en=body.text_en.strip(),
                    text_ml=ml or "(Malayalam not available: add it before approval.)"))
    db.commit()
    return {"ok": True, "translated": ml is not None and not body.text_ml}


@app.get("/api/translate/status")
def translate_status(p: Principal = Depends(current_user)):
    return {"enabled": translate.enabled(), "model": translate.MODEL, "disclaimer_en": translate.DISCLAIMER_EN,
            "disclaimer_ml": translate.DISCLAIMER_ML}


@app.post("/api/advisories/{aid}")
def review_advisory(aid: int, body: StatusBody, p: Principal = Depends(require("admin", "coordinator", "doctor", "asha")),
                    db: Session = Depends(get_db)):
    a = db.get(Advisory, aid)
    if not a or not p.can_see_site(a.site_code):
        raise HTTPException(404, "Not found.")
    if body.status in ("approved", "rejected") and p.role not in ("doctor", "admin"):
        raise HTTPException(403, "Only a doctor can approve or reject an advisory.")
    if body.status == "sent" and a.status != "approved":
        raise HTTPException(409, "An advisory must be approved by a doctor before it is sent.")
    if body.status not in ("approved", "rejected", "sent", "draft"):
        raise HTTPException(422, "Unknown status.")
    a.status = body.status
    if body.status in ("approved", "rejected"):
        a.reviewed_by = p.username
    db.commit()
    return {"ok": True}


@app.get("/api/advisories/send-list.csv")
def send_list(include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "asha")),
              db: Session = Depends(get_db)):
    """Approved household advisories with phone numbers, for sending by WhatsApp or SMS."""
    codes = sites_for(db, p, include_test)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["advisory_id", "site", "hh_id", "person_id", "name", "phone", "message_ml", "message_en", "note"])
    for a in db.query(Advisory).filter(Advisory.site_code.in_(codes), Advisory.status == "approved"):
        persons = db.query(Person).filter(Person.hh_id == a.hh_id).all() if a.hh_id else []
        for person in persons:
            pii = db.get(PersonPII, person.person_id)
            if pii and pii.phone and pii.contact_ok:
                w.writerow([a.id, a.site_code, a.hh_id, person.person_id, pii.name, pii.phone, a.text_ml, a.text_en,
                             translate.DISCLAIMER_EN])
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=send-list.csv"})


@app.get("/api/queries")
def list_queries(p: Principal = Depends(current_user)):
    return [{"name": k, "question": v} for k, v in rules.QUERIES.items()]


@app.get("/api/queries/{name}")
def run_query(name: str, include_test: bool = False, p: Principal = Depends(require(*ANALYSTS, "asha")),
              db: Session = Depends(get_db)):
    try:
        return rules.run_query(db, name, sites_for(db, p, include_test), p.pii)
    except KeyError:
        raise HTTPException(404, "Unknown question.")


# ------------------------------------------------------------------ records and exports (de-identified)

def _round(v, nd=3):
    return None if v is None else round(v, nd)


DATASETS = {
    "samples": lambda db, codes: [
        {"sample_id": s.sample_id, "site": s.site_code, "hh_id": s.hh_id, "point_type": s.point_type,
         "source": s.source, "taken_at": s.taken_at, "collected_on": s.collected_on, "lat": _round(s.lat),
         "lon": _round(s.lon), "tds": s.tds, "water_temp": s.water_temp, "rain_48h": s.rain_48h,
         "flooded_7d": s.flooded_7d, "aquadip": labs.get("aquadip").positive if "aquadip" in labs else None,
         "nio_lab": labs.get("nio_lab").positive if "nio_lab" in labs else None,
         "ecoli_cfu": labs["nio_lab"].ecoli_cfu if "nio_lab" in labs else None,
         "vibrio": labs["nio_lab"].vibrio if "nio_lab" in labs else None, "result": pos}
        for s, pos, labs in analytics.sample_results(db, codes)],
    "persons": lambda db, codes: [
        {"person_id": x.person_id, "hh_id": x.hh_id, "site": x.site_code, "sex": x.sex, "age": x.age,
         "resident": x.resident, "registered_on": x.registered_on}
        for x in db.query(Person).filter(Person.site_code.in_(codes), Person.consent_research.is_(True))],
    "households": lambda db, codes: [
        {c.name: getattr(h, c.name) for c in HouseholdSurvey.__table__.columns if c.name != "instance_id"}
        for h in db.query(HouseholdSurvey).filter(HouseholdSurvey.site_code.in_(codes))],
    "screening": lambda db, codes: [
        {c.name: getattr(r, c.name) for c in StationRecord.__table__.columns if c.name not in ("instance_id", "data")}
        for r in db.query(StationRecord).join(Person, Person.person_id == StationRecord.person_id)
        .filter(StationRecord.site_code.in_(codes), Person.consent_research.is_(True))],
    "visits": lambda db, codes: [
        {c.name: getattr(v, c.name) for c in HomeVisit.__table__.columns if c.name != "instance_id"}
        for v in db.query(HomeVisit).filter(HomeVisit.site_code.in_(codes))],
}


@app.get("/api/records/{dataset}")
def records(dataset: str, include_test: bool = False, site_type: str | None = None, limit: int = Query(200, le=5000),
            p: Principal = Depends(require(*ANALYSTS)), db: Session = Depends(get_db)):
    if dataset not in DATASETS:
        raise HTTPException(404, "Unknown dataset.")
    codes = sites_for(db, p, include_test)
    if site_type:
        codes = [c for c in codes if db.get(Site, c).site_type == site_type]
    rows = DATASETS[dataset](db, codes)
    return {"dataset": dataset, "total": len(rows), "rows": rows[:limit]}


@app.get("/api/export/{dataset}.csv")
def export(dataset: str, include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "scientist")),
           db: Session = Depends(get_db)):
    if dataset not in DATASETS:
        raise HTTPException(404, "Unknown dataset.")
    rows = DATASETS[dataset](db, sites_for(db, p, include_test))
    out = io.StringIO()
    if rows:
        w = csv.DictWriter(out, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename=observatory-{dataset}.csv"})


# ------------------------------------------------------------------ imports and admin

async def _read_csv(file: UploadFile) -> str:
    raw = await file.read()
    if len(raw) > 5_000_000:
        raise HTTPException(413, "File too large (5 MB maximum).")
    return raw.decode("utf-8-sig")


@app.post("/api/import/lab")
async def import_lab(file: UploadFile = File(...), p: Principal = Depends(require("admin", "coordinator", "scientist")),
                     db: Session = Depends(get_db)):
    """CSV columns: sample_id, method (nio_lab|aquadip), result (pos|neg), ecoli_cfu, coliform_cfu, vibrio, amr_note, tested_on."""
    text = await _read_csv(file)
    n, skipped = 0, []
    for i, row in enumerate(csv.DictReader(io.StringIO(text)), start=2):
        sid, method = (row.get("sample_id") or "").strip(), (row.get("method") or "nio_lab").strip()
        if not sid or method not in ("nio_lab", "aquadip") or row.get("result", "").strip() not in ("pos", "neg"):
            skipped.append(i)
            continue
        f = lambda k: float(row[k]) if row.get(k, "").strip() not in ("",) else None  # noqa: E731
        sync_odk.upsert_lab(db, sid, method, positive=row["result"].strip() == "pos", ecoli_cfu=f("ecoli_cfu"),
                            coliform_cfu=f("coliform_cfu"), vibrio=(row.get("vibrio") or None),
                            amr_note=(row.get("amr_note") or None),
                            tested_on=date.fromisoformat(row["tested_on"]) if row.get("tested_on") else None)
        n += 1
    db.commit()
    return {"imported": n, "skipped_rows": skipped}


@app.post("/api/import/tides")
async def import_tides(file: UploadFile = File(...), p: Principal = Depends(require(*MANAGERS, "scientist")),
                       db: Session = Depends(get_db)):
    n = weather.import_tides(db, await _read_csv(file))
    sources.record(db, "tides", True, detail=f"{n} days uploaded")
    return {"imported": n}


@app.post("/api/import/sites")
async def import_sites(file: UploadFile = File(...), p: Principal = Depends(require("admin")),
                       db: Session = Depends(get_db)):
    """The same sites.csv that is attached to the ODK forms."""
    n = 0
    for row in csv.DictReader(io.StringIO(await _read_csv(file))):
        code = row["name"].strip().upper()
        f = lambda k: float(row[k]) if row.get(k, "").strip() else None  # noqa: E731
        db.merge(Site(code=code, label=row["label"].strip(), site_type=row.get("site_type", "village").strip(),
                      lat=f("lat"), lon=f("lon"), active=True))
        n += 1
    db.commit()
    return {"sites": n}


@app.post("/api/sync")
def sync(p: Principal = Depends(require(*MANAGERS)), db: Session = Depends(get_db)):
    if config.FORM_SERVER == "odk" and not config.ODK_URL:
        raise HTTPException(503, "ODK Central is not configured (OBS_ODK_URL).")
    try:
        return {"result": sources.fetch(db, "odk")}
    except Exception as exc:
        raise HTTPException(502, f"Sync failed: {exc}")


@app.post("/api/weather/fetch")
def fetch_weather(p: Principal = Depends(require(*MANAGERS)), db: Session = Depends(get_db)):
    try:
        return {"result": sources.fetch(db, "open_meteo")}
    except Exception as exc:
        raise HTTPException(502, f"Weather fetch failed: {exc}")


# ------------------------------------------------------------------ data sources

@app.get("/api/admin/sources")
def data_sources(p: Principal = Depends(require(*MANAGERS)), db: Session = Depends(get_db)):
    return sources.status(db)


@app.post("/api/admin/sources/{key}/fetch")
def fetch_source(key: str, p: Principal = Depends(require(*MANAGERS)), db: Session = Depends(get_db)):
    if key not in {s["key"] for s in sources.SOURCES if s["fetch"]}:
        raise HTTPException(404, "This source cannot be fetched from here.")
    try:
        return {"ok": True, "result": sources.fetch(db, key)}
    except Exception as exc:
        return {"ok": False, "result": str(exc)}


# ------------------------------------------------------------------ household reports

@app.get("/api/reports")
def list_reports(include_test: bool = False, form_id: str | None = None, days: int = Query(90, ge=1, le=730),
                 p: Principal = Depends(current_user), db: Session = Depends(get_db)):
    return {"forms": [{"form_id": f.form_id, "title": f.title} for f in reports.visible_forms(db, p.role)],
            "rows": reports.list_reports(db, p, sites_for(db, p, include_test), form_id, days)}


@app.get("/api/stats/workdays")
def stats_workdays(include_test: bool = False, year: int | None = Query(None, ge=2020, le=2100),
                   p: Principal = Depends(require(*ANALYSTS, "ward")), db: Session = Depends(get_db)):
    return reports.workdays(db, sites_for(db, p, include_test), year,
                            daily_wage=reports.get_setting(db, "daily_wage_rs"))


class WageBody(BaseModel):
    daily_wage_rs: float | None = Field(default=None, ge=0, le=100000)


@app.post("/api/admin/settings/wage")
def set_wage(body: WageBody, p: Principal = Depends(require("admin")), db: Session = Depends(get_db)):
    reports.set_setting(db, "daily_wage_rs", body.daily_wage_rs)
    return {"ok": True}


@app.get("/api/admin/report-forms")
def report_forms(p: Principal = Depends(require("admin")), db: Session = Depends(get_db)):
    return [{"form_id": f.form_id, "title": f.title, "visible_roles": f.visible_roles.split(),
             "phone_groups": f.phone_groups, "active": f.active,
             "reports": db.query(Report).filter(Report.form_id == f.form_id).count()}
            for f in db.query(ReportForm).order_by(ReportForm.form_id)]


class ReportFormBody(BaseModel):
    title: str = Field(min_length=3, max_length=120)
    visible_roles: list[str]
    active: bool = True


@app.post("/api/admin/report-forms/{form_id}")
def update_report_form(form_id: str, body: ReportFormBody, p: Principal = Depends(require("admin")),
                       db: Session = Depends(get_db)):
    f = db.get(ReportForm, form_id)
    if not f:
        raise HTTPException(404, "Unknown report form.")
    bad = [r for r in body.visible_roles if r not in ROLES]
    if bad:
        raise HTTPException(422, f"Unknown role: {', '.join(bad)}")
    f.title, f.visible_roles, f.active = body.title.strip(), " ".join(sorted(set(body.visible_roles) | {"admin"})), body.active
    db.commit()
    return {"ok": True}


@app.get("/api/admin/report-form-template")
def report_form_template(p: Principal = Depends(require("admin"))):
    path = Path(__file__).resolve().parent.parent / "forms" / "dist" / "report_form_template.xlsx"
    return FileResponse(path, filename="report_form_template.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/admin/phone-groups")
def phone_groups(p: Principal = Depends(require("admin"))):
    if config.FORM_SERVER != "odk" or not config.ODK_URL:
        raise HTTPException(503, "ODK Central is not configured on the server.")
    try:
        return sync_odk.Central().app_users()
    except Exception as exc:
        raise HTTPException(502, f"Could not reach ODK Central: {exc}")


@app.post("/api/admin/report-forms")
async def upload_report_form(file: UploadFile = File(...), visible_roles: str = "admin coordinator doctor asha",
                             app_users: str = "", p: Principal = Depends(require("admin")),
                             db: Session = Depends(get_db)):
    """Publish a new or changed report form (an XLSForm) to ODK Central and register who may see its reports."""
    from openpyxl import load_workbook
    raw = await file.read()
    if len(raw) > 5_000_000:
        raise HTTPException(413, "File too large (5 MB maximum).")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True)
        rows = list(wb["settings"].iter_rows(values_only=True))
        settings = dict(zip(rows[0], rows[1]))
    except Exception:
        raise HTTPException(422, "This is not an XLSForm with a settings sheet.")
    form_id, title = str(settings.get("form_id") or "").strip(), str(settings.get("form_title") or "").strip()
    if not reports.is_report_form(form_id) or not form_id.replace("_", "").isalnum() or len(form_id) > 60:
        raise HTTPException(422, "The form_id in the settings sheet must start with obs_report_ and use only "
                                 "letters, numbers and underscores.")
    roles = [r for r in visible_roles.split() if r in ROLES]
    ids = [int(x) for x in app_users.split() if x.isdigit()]
    if config.FORM_SERVER == "odk" and config.ODK_URL:
        try:
            central = sync_odk.Central()
            central.publish_form(form_id, raw, ids, datetime.now(timezone.utc).strftime("%Y%m%d%H%M"))
            names = {u["id"]: u["name"] for u in central.app_users()}
        except Exception as exc:
            raise HTTPException(502, f"ODK Central: {exc}")
    else:
        raise HTTPException(503, "ODK Central is not configured on the server.")
    f = db.get(ReportForm, form_id) or ReportForm(form_id=form_id)
    f.title = title or form_id
    f.visible_roles = " ".join(sorted(set(roles) | {"admin"}))
    f.phone_groups = ", ".join(names.get(i, str(i)) for i in ids)
    f.active = True
    db.merge(f)
    db.commit()
    return {"ok": True, "form_id": form_id, "title": f.title}


class NewUser(BaseModel):
    username: str = Field(min_length=3, max_length=60, pattern=r"^[a-z0-9._-]+$")
    full_name: str
    role: str
    sites: str = "*"
    password: str = Field(min_length=10)


@app.post("/api/admin/users")
def create_user(body: NewUser, p: Principal = Depends(require("admin")), db: Session = Depends(get_db)):
    if body.role not in ROLES:
        raise HTTPException(422, f"Role must be one of {', '.join(ROLES)}.")
    db.merge(User(username=body.username, full_name=body.full_name, role=body.role, sites=body.sites,
                  password_hash=hash_password(body.password), active=True))
    db.commit()
    return {"ok": True}


@app.get("/api/admin/audit")
def audit(limit: int = Query(200, le=2000), p: Principal = Depends(require("admin")), db: Session = Depends(get_db)):
    rows = db.query(AuditLog).order_by(AuditLog.id.desc()).limit(limit)
    return [{"at": a.at.isoformat() if a.at else None, "user": a.username, "method": a.method, "path": a.path,
             "status": a.status, "ip": a.ip} for a in rows]


# ------------------------------------------------------------------ IDSP module

def _week(week: str | None) -> date:
    """The IDSP week (Monday to Sunday) holding `week`, or the last complete week."""
    if week:
        try:
            return idsp.week_start(date.fromisoformat(week))
        except ValueError:
            raise HTTPException(422, "week must be a date, YYYY-MM-DD.")
    return idsp.week_start(date.today()) - timedelta(days=7)


def _idsp_case(db: Session, p: Principal, cid: int) -> IdspCase:
    c = db.get(IdspCase, cid)
    if not c or not p.can_see_site(c.site_code):
        raise HTTPException(404, "No such case.")
    return c


@app.get("/api/idsp/meta")
def idsp_meta(p: Principal = Depends(require(*IDSP_VIEW))):
    return idsp.meta()


@app.get("/api/idsp/overview")
def idsp_overview(week: str | None = None, include_test: bool = False, p: Principal = Depends(require(*IDSP_VIEW)),
                  db: Session = Depends(get_db)):
    return idsp.overview(db, sites_for(db, p, include_test), _week(week))


@app.get("/api/idsp/queue")
def idsp_queue(include_test: bool = False, p: Principal = Depends(require(*IDSP_VIEW)), db: Session = Depends(get_db)):
    if p.role in ("scientist", "lab"):
        raise HTTPException(403, "Your role cannot open the verification queue.")
    return idsp.queue(db, sites_for(db, p, include_test), p.pii)


@app.get("/api/idsp/cases")
def idsp_cases(week: str | None = None, status: str | None = None, include_test: bool = False,
               p: Principal = Depends(require(*IDSP_VIEW)), db: Session = Depends(get_db)):
    if status and status not in idsp.STATUS_LABEL:
        raise HTTPException(422, "Unknown status.")
    return idsp.cases(db, sites_for(db, p, include_test), p.pii, _week(week) if week else None, status)


class VerifyBody(BaseModel):
    outcome: str
    note: str | None = Field(default=None, max_length=300)
    duplicate_of: int | None = None


@app.post("/api/idsp/cases/{cid}/verify")
def idsp_verify(cid: int, body: VerifyBody, p: Principal = Depends(require(*IDSP_VERIFY)),
                db: Session = Depends(get_db)):
    c = _idsp_case(db, p, cid)
    if body.duplicate_of:
        _idsp_case(db, p, body.duplicate_of)
    try:
        idsp.verify(db, c, body.outcome, p.username, p.role, body.note, body.duplicate_of)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    return {"ok": True}


class PresumptiveBody(BaseModel):
    diagnosis: str


@app.post("/api/idsp/cases/{cid}/presumptive")
def idsp_presumptive(cid: int, body: PresumptiveBody, p: Principal = Depends(require("admin", "mo")),
                     db: Session = Depends(get_db)):
    try:
        idsp.set_presumptive(db, _idsp_case(db, p, cid), body.diagnosis, p.username)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    return {"ok": True}


class LabBody(BaseModel):
    test: str
    result: str
    sample_on: date | None = None
    sample_id: str | None = Field(default=None, max_length=40)


@app.post("/api/idsp/cases/{cid}/lab")
def idsp_lab(cid: int, body: LabBody, p: Principal = Depends(require("admin", "mo", "lab")),
             db: Session = Depends(get_db)):
    c = _idsp_case(db, p, cid)
    if c.status != "verified":
        raise HTTPException(409, "Verify the case before adding a laboratory result.")
    try:
        idsp.set_lab(db, c, body.test, body.result, p.username, body.sample_on, body.sample_id)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    return {"ok": True}


@app.get("/api/idsp/weekly")
def idsp_weekly(week: str | None = None, include_test: bool = False, p: Principal = Depends(require(*IDSP_VIEW)),
                db: Session = Depends(get_db)):
    return idsp.weekly(db, sites_for(db, p, include_test), _week(week))


def _csv(rows, filename):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})


@app.get("/api/idsp/weekly.csv")
def idsp_weekly_csv(week: str | None = None, include_test: bool = False, p: Principal = Depends(require(*IDSP_EXPORT)),
                    db: Session = Depends(get_db)):
    data = idsp.weekly(db, sites_for(db, p, include_test), _week(week))
    return _csv(idsp.weekly_csv_rows(data), f"idsp-form-s-p-{data['week_start']}.csv")


@app.get("/api/idsp/linelist.csv")
def idsp_linelist_csv(week: str | None = None, include_test: bool = False,
                      p: Principal = Depends(require(*IDSP_EXPORT)), db: Session = Depends(get_db)):
    ws = _week(week)
    rows = idsp.cases(db, sites_for(db, p, include_test), p.pii, ws, "verified")
    return _csv(idsp.linelist_csv_rows(rows, p.pii), f"idsp-line-list-{ws.isoformat()}.csv")


@app.get("/api/idsp/alerts")
def idsp_alerts(include_test: bool = False, p: Principal = Depends(require(*IDSP_VIEW)), db: Session = Depends(get_db)):
    return idsp.alert_rows(db, sites_for(db, p, include_test))


class AlertAction(BaseModel):
    action: str
    note: str | None = Field(default=None, max_length=500)


@app.post("/api/idsp/alerts/{aid}")
def idsp_alert_act(aid: int, body: AlertAction, p: Principal = Depends(require(*IDSP_ALERT)),
                   db: Session = Depends(get_db)):
    a = db.get(IdspAlert, aid)
    if not a or not p.can_see_site(a.site_code):
        raise HTTPException(404, "No such alert.")
    try:
        idsp.act_on_alert(db, a, body.action, p.username, body.note)
    except ValueError as exc:
        raise HTTPException(422, str(exc))
    return {"ok": True}


@app.post("/api/idsp/run")
def idsp_run(include_test: bool = False, p: Principal = Depends(require("admin", "coordinator", "mo", "dso")),
             db: Session = Depends(get_db)):
    """Capture new community signals and check the alert thresholds now (also runs every hour)."""
    captured = idsp.capture_signals(db)
    return {"signals_captured": captured, **idsp.generate_alerts(db, sites_for(db, p, include_test))}


@app.get("/api/idsp/thresholds")
def idsp_thresholds(p: Principal = Depends(require(*IDSP_VIEW)), db: Session = Depends(get_db)):
    return idsp.thresholds(db)


@app.post("/api/idsp/thresholds")
def idsp_set_thresholds(body: dict, p: Principal = Depends(require("admin", "dso")), db: Session = Depends(get_db)):
    clean = {}
    for k, v in body.items():
        if k not in idsp.TH_DEFAULT:
            raise HTTPException(422, f"Unknown setting: {k}")
        try:
            num = float(v)
        except (TypeError, ValueError):
            raise HTTPException(422, f"{k} must be a number.")
        if not math.isfinite(num) or num <= 0 or num > 10000:
            raise HTTPException(422, f"{k} is out of range.")
        clean[k] = int(num) if num == int(num) else num
    idsp.set_thresholds(db, clean)
    return {"ok": True, "thresholds": idsp.thresholds(db)}


# ------------------------------------------------------------------ dashboard

WEB = Path(__file__).resolve().parent.parent / "web"
app.mount("/static", StaticFiles(directory=WEB), name="static")


@app.get("/")
def index():
    return FileResponse(WEB / "index.html")


@app.get("/install")
def install_page():
    """Public page: how to install and set up the field app. Holds no setup codes."""
    return FileResponse(WEB / "install.html")


@app.exception_handler(Exception)
async def unhandled(request: Request, exc: Exception):
    return JSONResponse({"detail": "Something went wrong on the server."}, status_code=500)
