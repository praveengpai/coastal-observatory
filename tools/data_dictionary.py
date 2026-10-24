"""Generate docs/data_dictionary.md from the XLSForms (the single source of truth for field names).

    python -m tools.data_dictionary
"""
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SKIP = {"begin_group", "end_group", "note", "start", "end", "today", "deviceid"}


def main():
    out = ["# Data dictionary", "",
           "Generated from the ODK forms in `forms/dist`. Field names here are the column names in exports "
           "and the keys stored in the database. Do not rename a field once data has been collected.", ""]
    for path in sorted((ROOT / "forms" / "dist").glob("*.xlsx")):
        ws = load_workbook(path, read_only=True)["survey"]
        rows = list(ws.iter_rows(values_only=True))
        head = rows[0]
        ix = {h: i for i, h in enumerate(head)}
        out += [f"## {path.stem}", "", "| Field | Type | Question (English) | Shown when | Calculation |",
                "| --- | --- | --- | --- | --- |"]
        for r in rows[1:]:
            t = r[ix["type"]]
            if not t or t in SKIP:
                continue
            cell = lambda k: (r[ix[k]] or "").replace("|", "\\|").replace("\n", " ")  # noqa: E731
            out.append(f"| `{r[ix['name']]}` | {t} | {cell('label::English (en)')} | {cell('relevant')} | {cell('calculation')} |")
        out.append("")
    (ROOT / "docs" / "data_dictionary.md").write_text("\n".join(out), encoding="utf-8")
    print("docs/data_dictionary.md written")


if __name__ == "__main__":
    main()
