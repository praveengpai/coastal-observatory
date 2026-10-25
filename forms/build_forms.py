"""Builds the Observatory's ODK XLSForms.

obs_camp_registration : welcome desk + waiting-area questionnaire; adds each person to the phones' household list
obs_camp_station      : measurements, eye, dental, geriatric and doctor stations
obs_household_visit   : ASHA and volunteer visits between camps
obs_water_sample      : water samples and Aquadip results
obs_report_*          : household reports (scan the door sticker, pick or add the person, report);
                        report_form_template.xlsx is the starting point for new report forms
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table

OUT = sys.argv[1]

SURVEY_COLS = [
    "type", "name", "label::English (en)", "label::Malayalam (ml)",
    "hint::English (en)", "hint::Malayalam (ml)", "required", "relevant",
    "constraint", "constraint_message::English (en)", "constraint_message::Malayalam (ml)",
    "calculation", "default", "appearance", "parameters", "choice_filter", "save_to",
]
CHOICE_COLS = ["list_name", "name", "label::English (en)", "label::Malayalam (ml)", "site_type"]

CHECK_EN = "Please check this value."
CHECK_ML = "ഈ മൂല്യം ഒന്നുകൂടി പരിശോധിക്കുക."
VOICE_EN = "Tip: tap the microphone key on the keyboard and speak, in Malayalam or English."
VOICE_ML = "കീബോർഡിലെ മൈക്ക് അമർത്തി മലയാളത്തിലോ ഇംഗ്ലീഷിലോ പറഞ്ഞാൽ മതി."


def q(t, name, en="", ml=None, hen="", hml="", req="", rel="", con="", cmsg_en="", cmsg_ml="",
      calc="", default="", app="", params="", cf="", save=""):
    if ml is None:
        ml = en
    if con and not cmsg_en:
        cmsg_en, cmsg_ml = CHECK_EN, CHECK_ML
    return [t, name, en, ml, hen, hml, req, rel, con, cmsg_en, cmsg_ml, calc, default, app, params, cf, save]


def c(lst, name, en, ml=None, site_type=""):
    return [lst, name, en, en if ml is None else ml, site_type]


NONE_EXCL = "not(selected(., 'none') and count-selected(.) > 1)"
NONE_MSG_EN = "'None' cannot be combined with other answers."
NONE_MSG_ML = "'ഒന്നുമില്ല' മറ്റ് ഉത്തരങ്ങളോടൊപ്പം തിരഞ്ഞെടുക്കാൻ പാടില്ല."

PERSON_RE = "regex(., '^OBS-P-[0-9]{6}$')"
HH_RE = "regex(., '^[A-Z]{2,4}-[0-9]{3}$')"

# ---------------------------------------------------------------- shared blocks

def camp_block(role_list, role_en="Your role today", role_ml="ഇന്ന് നിങ്ങളുടെ ചുമതല"):
    return [
        q("start", "start"), q("end", "end"), q("today", "today"), q("deviceid", "deviceid"),
        q("start-geopoint", "camp_gps"),
        q("begin_group", "camp", "Camp details", "ക്യാമ്പ് വിവരങ്ങൾ", app="field-list"),
        q("note", "camp_note",
          "Camp details carry over from your last form. Change them only if you have moved to another site. "
          "(Malayalam text is a machine-assisted translation for accessibility; English is the reference.)",
          "ക്യാമ്പ് വിവരങ്ങൾ കഴിഞ്ഞ ഫോമിൽ നിന്ന് തനിയെ വരും. സ്ഥലം മാറിയാൽ മാത്രം മാറ്റുക. "
          "(മലയാളം യന്ത്രസഹായത്തോടെയുള്ള വിവർത്തനമാണ്; ഇംഗ്ലീഷ് പാഠമാണ് ആധികാരികം.)"),
        q("select_one_from_file sites.csv", "camp_site", "Site", "സ്ഥലം", req="yes",
          default="${last-saved#camp_site}"),
        q(f"select_one {role_list}", "role", role_en, role_ml, req="yes",
          default="${last-saved#role}"),
        q("end_group", "camp"),
        q("calculate", "site_type", calc="instance('sites')/root/item[name=${camp_site}]/site_type"),
        q("calculate", "camp_id", calc="concat(${camp_site}, '-', format-date(today(), '%Y%m%d'))"),
    ]


def person_scan_block(label_en, label_ml, save=False):
    return [
        q("begin_group", "pid", label_en, label_ml, app="field-list"),
        q("select_one id_method", "pid_method", "Person ID: scan or type?", "വ്യക്തിയുടെ ഐഡി: സ്കാൻ ചെയ്യണോ ടൈപ്പ് ചെയ്യണോ?",
          req="yes", default="scan"),
        q("barcode", "pid_scan", "Scan the person ID sticker on the health card",
          "ഹെൽത്ത് കാർഡിലെ ഐഡി സ്റ്റിക്കർ സ്കാൻ ചെയ്യുക", req="yes", rel="${pid_method}='scan'",
          con=PERSON_RE, cmsg_en="This is not a person ID sticker (format OBS-P-000123).",
          cmsg_ml="ഇത് വ്യക്തിയുടെ ഐഡി സ്റ്റിക്കർ അല്ല (ഉദാ: OBS-P-000123)."),
        q("text", "pid_type", "Type the person ID", "വ്യക്തിയുടെ ഐഡി ടൈപ്പ് ചെയ്യുക", hen="Format: OBS-P-000123",
          hml="ഉദാ: OBS-P-000123", req="yes", rel="${pid_method}='type'", con=PERSON_RE,
          cmsg_en="Use the format OBS-P-000123.", cmsg_ml="OBS-P-000123 എന്ന രൂപത്തിൽ നൽകുക."),
        q("end_group", "pid"),
        q("calculate", "person_id", calc="if(${pid_method}='scan', ${pid_scan}, ${pid_type})",
          save="person_id" if save else ""),
    ]


# ---------------------------------------------------------------- shared choices

def common_choices():
    rows = []
    rows += [
        c("yes_no", "yes", "Yes", "അതെ"), c("yes_no", "no", "No", "ഇല്ല"),
        c("ynd", "yes", "Yes", "അതെ"), c("ynd", "no", "No", "ഇല്ല"), c("ynd", "dk", "Don't know", "അറിയില്ല"),
        c("id_method", "scan", "Scan", "സ്കാൻ"), c("id_method", "type", "Type", "ടൈപ്പ്"),
    ]
    return rows


# ================================================================ FORM 1

def form_registration():
    s = camp_block("reg_role")
    s += [
        # The household member list on every phone (an ODK entity list). Registration adds each new person once.
        q("csv-external", "persons"),
        q("begin_group", "consent", "Consent", "സമ്മതം", app="field-list"),
        q("note", "consent_care",
          "[Replace with ethics-approved text] Check-ups and medicines at this camp are free for everyone.",
          "[അംഗീകരിച്ച സമ്മതപത്രം ഇവിടെ ചേർക്കുക] ഈ ക്യാമ്പിലെ പരിശോധനയും മരുന്നും എല്ലാവർക്കും സൗജന്യമാണ്."),
        q("select_one yes_no", "consent_research",
          "May we use this information, without names, for the water and health study? Care is the same whether or not you agree.",
          "പേര് ഒഴിവാക്കി ഈ വിവരങ്ങൾ ജല-ആരോഗ്യ പഠനത്തിനായി ഉപയോഗിക്കാൻ സമ്മതമാണോ? സമ്മതിച്ചാലും ഇല്ലെങ്കിലും ചികിത്സ ഒരുപോലെ ലഭിക്കും.",
          req="yes"),
        q("end_group", "consent"),
    ]
    s += person_scan_block("Person ID", "വ്യക്തിയുടെ ഐഡി", save=True)
    s += [
        q("begin_group", "hid", "Household ID", "വീടിന്റെ ഐഡി", app="field-list"),
        q("select_one id_method", "hh_method", "Household ID: scan or type?", "വീടിന്റെ ഐഡി: സ്കാൻ ചെയ്യണോ ടൈപ്പ് ചെയ്യണോ?",
          req="yes", default="scan"),
        q("barcode", "hh_scan", "Scan the household sticker", "വീടിന്റെ സ്റ്റിക്കർ സ്കാൻ ചെയ്യുക", req="yes",
          rel="${hh_method}='scan'", con=HH_RE, cmsg_en="This is not a household sticker (format CVB-014).",
          cmsg_ml="ഇത് വീടിന്റെ സ്റ്റിക്കർ അല്ല (ഉദാ: CVB-014)."),
        q("text", "hh_type", "Type the household ID", "വീടിന്റെ ഐഡി ടൈപ്പ് ചെയ്യുക", hen="Format: CVB-014",
          hml="ഉദാ: CVB-014", req="yes", rel="${hh_method}='type'", con=HH_RE,
          cmsg_en="Use the format CVB-014.", cmsg_ml="CVB-014 എന്ന രൂപത്തിൽ നൽകുക."),
        q("end_group", "hid"),
        q("calculate", "hh_id", calc="if(${hh_method}='scan', ${hh_scan}, ${hh_type})", save="hh_id"),

        q("begin_group", "person", "About the person", "വ്യക്തിയെ കുറിച്ച്", app="field-list"),
        q("text", "name", "Name", "പേര്", hen="Manglish or Malayalam keyboard both work.",
          hml="മംഗ്ലീഷ് അല്ലെങ്കിൽ മലയാളം കീബോർഡ് ഉപയോഗിക്കാം.", req="yes", save="full_name"),
        q("integer", "age", "Age (years)", "പ്രായം (വർഷം)", req="yes", con=". >= 0 and . <= 120", save="age"),
        q("select_one sex", "sex", "Sex", "ലിംഗം", req="yes", app="horizontal-compact", save="sex"),
        q("text", "phone", "Phone number (for results)", "ഫോൺ നമ്പർ (ഫലം അറിയിക്കാൻ)", app="numbers",
          con="regex(., '^[6-9][0-9]{9}$')", cmsg_en="Enter a 10-digit mobile number.",
          cmsg_ml="10 അക്ക മൊബൈൽ നമ്പർ നൽകുക."),
        q("select_one yes_no", "resident", "Does this person live here permanently?",
          "ഈ വ്യക്തി ഇവിടെ സ്ഥിരതാമസമാണോ?", req="yes", default="yes", app="horizontal-compact"),
        q("end_group", "person"),

        q("begin_group", "visitor", "Visitor details", "സന്ദർശകന്റെ വിവരങ്ങൾ", rel="${resident}='no'", app="field-list"),
        q("text", "home_address", "Home address", "വീട്ടു വിലാസം", hen=VOICE_EN, hml=VOICE_ML, req="yes"),
        q("text", "home_lsg", "Panchayat / municipality / ward", "പഞ്ചായത്ത് / നഗരസഭ / വാർഡ്"),
        q("select_one stay", "stay", "How long staying here?", "ഇവിടെ എത്ര കാലമായി താമസിക്കുന്നു?"),
        q("end_group", "visitor"),

        q("select_one yes_no", "first_member",
          "Is this the first person from this household registered at this camp?",
          "ഈ ക്യാമ്പിൽ ഈ വീട്ടിൽ നിന്ന് ആദ്യം രജിസ്റ്റർ ചെയ്യുന്നത് ഇദ്ദേഹമാണോ?", req="yes",
          rel="${consent_research}='yes' and ${resident}='yes'", app="horizontal-compact"),
    ]
    # ---- Part A household
    HH = "${consent_research}='yes' and ${resident}='yes' and ${first_member}='yes'"
    s += [
        q("begin_group", "household", "Part A: Household and water", "ഭാഗം A: വീടും വെള്ളവും", rel=HH),
        q("begin_group", "hh_people", "People in the house", "വീട്ടിലെ ആളുകൾ", app="field-list"),
        q("integer", "hh_size", "Number of people living in this house", "ഈ വീട്ടിൽ താമസിക്കുന്നവരുടെ എണ്ണം",
          req="yes", con=". >= 1 and . <= 30"),
        q("integer", "hh_n60", "How many are aged 60 or above?", "ഇതിൽ 60 വയസ്സോ അതിനു മുകളിലോ ഉള്ളവർ എത്ര?",
          req="yes", con=". >= 0 and . <= ${hh_size}"),
        q("integer", "hh_n5", "How many children under 5?", "5 വയസ്സിൽ താഴെയുള്ള കുട്ടികൾ എത്ര?",
          req="yes", con=". >= 0 and . <= ${hh_size}"),
        q("end_group", "hh_people"),
        q("begin_group", "hh_water", "Drinking water", "കുടിവെള്ളം", app="field-list"),
        q("select_one water_src", "water_drink", "Main source of drinking water", "കുടിവെള്ളത്തിന്റെ പ്രധാന ഉറവിടം", req="yes"),
        q("select_one yes_no", "cook_same", "Same water used for cooking?", "പാചകത്തിനും ഇതേ വെള്ളമാണോ?", req="yes",
          app="horizontal-compact"),
        q("select_one water_src", "water_cook", "Water for cooking", "പാചകത്തിനുള്ള വെള്ളം", req="yes",
          rel="${cook_same}='no'"),
        q("select_multiple treat", "treat", "How is drinking water treated?", "കുടിവെള്ളം എങ്ങനെ ശുദ്ധീകരിക്കുന്നു?",
          req="yes", con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("select_one storage", "storage", "How is drinking water stored?", "കുടിവെള്ളം എങ്ങനെ സൂക്ഷിക്കുന്നു?", req="yes"),
        q("end_group", "hh_water"),
        q("begin_group", "hh_salt", "Salt and sanitation", "ഉപ്പുരസവും ശുചിത്വവും", app="field-list"),
        q("select_one freq3", "salty", "Does your well or tap water ever taste salty?",
          "കിണർ / പൈപ്പ് വെള്ളത്തിന് ഉപ്പുരസം തോന്നാറുണ്ടോ?", req="yes"),
        q("select_multiple season", "salty_season", "In which season?", "ഏത് കാലത്ത്?", rel="${salty}!='never'"),
        q("select_one toilet", "toilet", "Type of toilet", "ശൗചാലയത്തിന്റെ തരം", req="yes"),
        q("select_one dist", "septic_dist", "Distance from septic tank or pit to the well",
          "സെപ്റ്റിക് ടാങ്കിൽ / കുഴിയിൽ നിന്ന് കിണറിലേക്കുള്ള ദൂരം",
          rel="selected(${water_drink}, 'own_well') or selected(${water_drink}, 'shared_well')"),
        q("end_group", "hh_salt"),
        q("begin_group", "hh_flood", "Flooding and animals", "വെള്ളപ്പൊക്കവും മൃഗങ്ങളും", app="field-list"),
        q("select_one yes_no", "flood", "In the past 12 months, did water enter the house or compound?",
          "കഴിഞ്ഞ 12 മാസത്തിൽ വെള്ളം വീട്ടിലോ മുറ്റത്തോ കയറിയോ?", req="yes", app="horizontal-compact"),
        q("integer", "flood_times", "How many times?", "എത്ര തവണ?", req="yes", rel="${flood}='yes'",
          con=". >= 1 and . <= 50"),
        q("select_multiple animals", "animals", "Animals or birds kept at home", "വീട്ടിൽ വളർത്തുന്ന മൃഗങ്ങൾ / പക്ഷികൾ",
          req="yes", con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("select_one freq3", "rats", "Do you see rats around the house or water storage?",
          "വീടിനു ചുറ്റും / വെള്ളം സൂക്ഷിക്കുന്നിടത്ത് എലികളെ കാണാറുണ്ടോ?", req="yes"),
        q("end_group", "hh_flood"),
        q("begin_group", "hh_flats", "Flats: sump and tank", "ഫ്ലാറ്റ്: സംപും ടാങ്കും", rel="${site_type}='flats'",
          app="field-list"),
        q("select_one sump_mgr", "sump_mgr", "Who manages the sump and tank?", "സംപും ടാങ്കും ആരാണ് നോക്കുന്നത്?"),
        q("select_one last_clean", "tank_clean", "When was the tank last cleaned?", "ടാങ്ക് അവസാനമായി വൃത്തിയാക്കിയത് എപ്പോൾ?"),
        q("end_group", "hh_flats"),
        q("begin_group", "hh_tds", "Water salt test (TDS)", "വെള്ളത്തിലെ ഉപ്പ് പരിശോധന (TDS)", app="field-list"),
        q("select_one yes_no", "bottle", "Did they bring a bottle of their drinking water?",
          "കുടിവെള്ളം കുപ്പിയിൽ കൊണ്ടുവന്നിട്ടുണ്ടോ?", req="yes", app="horizontal-compact"),
        q("integer", "tds", "TDS meter reading (mg/L)", "TDS മീറ്റർ റീഡിംഗ് (mg/L)", req="yes", rel="${bottle}='yes'",
          con=". >= 0 and . <= 20000"),
        q("note", "tds_flag", "Reading above 1,000 mg/L: tell the person their water tastes salty for a reason, and mark the health card.",
          "1000 mg/L-ൽ കൂടുതൽ: വെള്ളത്തിൽ ഉപ്പ് കൂടുതലാണെന്ന് പറയുക, ഹെൽത്ത് കാർഡിൽ രേഖപ്പെടുത്തുക.",
          rel="${tds} > 1000"),
        q("end_group", "hh_tds"),
        q("end_group", "household"),
    ]
    # ---- Part B individual (research)
    R = "${consent_research}='yes'"
    s += [
        q("begin_group", "indiv", "Part B: Recent illness", "ഭാഗം B: അടുത്തിടെയുള്ള അസുഖങ്ങൾ", rel=R),
        q("begin_group", "occ_g", "Work and exposure", "ജോലിയും സമ്പർക്കവും", app="field-list"),
        q("select_one occ", "occupation", "Main occupation", "പ്രധാന തൊഴിൽ", req="yes", rel="${age} >= 15"),
        q("select_one yes_no", "wade", "In the past month, walked through flood or canal water?",
          "കഴിഞ്ഞ ഒരു മാസത്തിൽ വെള്ളക്കെട്ടിലോ കനാൽ വെള്ളത്തിലോ ഇറങ്ങി നടന്നിട്ടുണ്ടോ?", req="yes",
          app="horizontal-compact"),
        q("end_group", "occ_g"),
        q("begin_group", "sx_g", "Past 2 weeks", "കഴിഞ്ഞ 2 ആഴ്ച", app="field-list"),
        q("select_multiple sx", "sx", "In the past 2 weeks, any of these?", "കഴിഞ്ഞ 2 ആഴ്ചയിൽ ഇവയിൽ ഏതെങ്കിലും ഉണ്ടായോ?",
          req="yes", con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("integer", "sx_days", "Days unwell", "എത്ര ദിവസം അസുഖമുണ്ടായിരുന്നു?", req="yes",
          rel="count-selected(${sx}) > 0 and not(selected(${sx}, 'none'))", con=". >= 1 and . <= 14"),
        q("select_one sx_care", "sx_care", "Treatment taken", "എടുത്ത ചികിത്സ", req="yes",
          rel="count-selected(${sx}) > 0 and not(selected(${sx}, 'none'))"),
        q("select_one yes_no", "sx_missed", "Missed work or school?", "ജോലി / സ്കൂൾ മുടങ്ങിയോ?",
          rel="count-selected(${sx}) > 0 and not(selected(${sx}, 'none'))", app="horizontal-compact"),
        q("end_group", "sx_g"),
        q("end_group", "indiv"),
    ]
    # ---- health history (service)
    s += [
        q("begin_group", "hist", "Health history", "ആരോഗ്യ ചരിത്രം", app="field-list"),
        q("select_multiple cond", "conditions", "Known illnesses", "അറിയാവുന്ന രോഗങ്ങൾ", req="yes",
          con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("text", "meds", "Medicines taken now", "ഇപ്പോൾ കഴിക്കുന്ന മരുന്നുകൾ", hen=VOICE_EN, hml=VOICE_ML,
          rel="not(selected(${conditions}, 'none'))", app="multiline"),
        q("select_one tob", "tobacco", "Tobacco, paan or areca use", "പുകയില / മുറുക്ക് / പാക്ക് ഉപയോഗം", req="yes",
          rel="${age} >= 12"),
        q("end_group", "hist"),
        q("select_one yes_no", "contact_ok", "May we phone you with your results?", "ഫലങ്ങൾ ഫോണിൽ അറിയിക്കാമോ?",
          req="yes", rel="string-length(${phone}) > 0", app="horizontal-compact"),
        q("note", "done_note",
          "Done. Write ${person_id} and ${hh_id} on the health card and send the person to measurements.",
          "പൂർത്തിയായി. ${person_id}, ${hh_id} ഹെൽത്ത് കാർഡിൽ എഴുതി അളവെടുക്കലിലേക്ക് അയയ്ക്കുക."),
    ]

    ch = common_choices()
    ch += [c("reg_role", n, en, ml) for n, en, ml in [
        ("volunteer", "Volunteer", "വോളണ്ടിയർ"), ("asha", "ASHA", "ആശാ പ്രവർത്തക"),
        ("oh_vol", "One Health volunteer", "വൺ ഹെൽത്ത് വോളണ്ടിയർ"), ("msw", "MSW", "എം.എസ്.ഡബ്ല്യു"),
        ("coord", "Coordinator", "കോർഡിനേറ്റർ")]]
    ch += [c("sex", n, en, ml) for n, en, ml in [("m", "Male", "പുരുഷൻ"), ("f", "Female", "സ്ത്രീ"), ("o", "Other", "മറ്റുള്ളവ")]]
    ch += [c("stay", n, en, ml) for n, en, ml in [
        ("lt1m", "Less than a month", "ഒരു മാസത്തിൽ താഴെ"), ("1to6m", "1–6 months", "1–6 മാസം"),
        ("gt6m", "More than 6 months", "6 മാസത്തിലധികം"), ("visit", "Only visiting today", "ഇന്ന് സന്ദർശനം മാത്രം")]]
    ch += [c("water_src", n, en, ml) for n, en, ml in [
        ("own_well", "Own well", "സ്വന്തം കിണർ"), ("shared_well", "Shared or public well", "പൊതു / പങ്കിട്ട കിണർ"),
        ("tap", "Tap (KWA / piped)", "പൈപ്പ് വെള്ളം (KWA)"), ("tanker", "Tanker", "ടാങ്കർ വെള്ളം"),
        ("borewell", "Borewell", "കുഴൽക്കിണർ"), ("packaged", "Bottled or can water", "കുപ്പി / കാൻ വെള്ളം"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("treat", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("boil", "Boil", "തിളപ്പിക്കും"), ("filter", "Cloth or candle filter", "ഫിൽട്ടർ"),
        ("ro", "RO / purifier", "ആർ.ഒ. / പ്യൂരിഫയർ"), ("chlorine", "Chlorine", "ക്ലോറിൻ")]]
    ch += [c("storage", n, en, ml) for n, en, ml in [
        ("covered", "Covered pot or can", "അടച്ച പാത്രം"), ("open", "Open container", "തുറന്ന പാത്രം"),
        ("overhead", "Overhead tank", "മുകളിലെ ടാങ്ക്"), ("sump", "Sump", "സംപ് (ഭൂഗർഭ ടാങ്ക്)"),
        ("none", "Not stored", "സൂക്ഷിക്കാറില്ല")]]
    ch += [c("freq3", n, en, ml) for n, en, ml in [
        ("never", "Never", "ഒരിക്കലുമില്ല"), ("sometimes", "Sometimes", "ചിലപ്പോൾ"), ("often", "Often", "പലപ്പോഴും")]]
    ch += [c("season", n, en, ml) for n, en, ml in [
        ("summer", "Summer (Mar–May)", "വേനൽ (മാർച്ച്–മേയ്)"), ("monsoon", "Monsoon (Jun–Sep)", "മഴക്കാലം (ജൂൺ–സെപ്റ്റംബർ)"),
        ("post", "After monsoon (Oct–Feb)", "മഴയ്ക്ക് ശേഷം (ഒക്ടോബർ–ഫെബ്രുവരി)")]]
    ch += [c("toilet", n, en, ml) for n, en, ml in [
        ("septic", "Septic tank", "സെപ്റ്റിക് ടാങ്ക്"), ("pit", "Pit latrine", "കുഴി കക്കൂസ്"),
        ("sewer", "Sewer connection", "സീവേജ് കണക്ഷൻ"), ("shared", "Shared or public toilet", "പൊതു ശൗചാലയം"),
        ("none", "No toilet", "ശൗചാലയം ഇല്ല")]]
    ch += [c("dist", n, en, ml) for n, en, ml in [
        ("lt10", "Under 10 m", "10 മീറ്ററിൽ താഴെ"), ("10to30", "10–30 m", "10–30 മീറ്റർ"),
        ("gt30", "Over 30 m", "30 മീറ്ററിൽ കൂടുതൽ"), ("dk", "Don't know", "അറിയില്ല")]]
    ch += [c("animals", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("dog", "Dog", "നായ"), ("cat", "Cat", "പൂച്ച"),
        ("cattle", "Cattle", "കന്നുകാലികൾ"), ("goat", "Goat", "ആട്"), ("poultry", "Hens or ducks", "കോഴി / താറാവ്")]]
    ch += [c("sump_mgr", n, en, ml) for n, en, ml in [
        ("rwa", "Residents' association", "റസിഡന്റ്സ് അസോസിയേഷൻ"), ("agency", "Builder or agency", "ബിൽഡർ / ഏജൻസി"),
        ("each", "Each flat", "ഓരോ ഫ്ലാറ്റും"), ("dk", "Don't know", "അറിയില്ല")]]
    ch += [c("last_clean", n, en, ml) for n, en, ml in [
        ("lt6m", "Within 6 months", "6 മാസത്തിനുള്ളിൽ"), ("6to12", "6–12 months ago", "6–12 മാസം മുമ്പ്"),
        ("gt1y", "Over a year ago", "ഒരു വർഷത്തിലേറെ മുമ്പ്"), ("dk", "Don't know", "അറിയില്ല")]]
    ch += [c("occ", n, en, ml) for n, en, ml in [
        ("fishing", "Fishing", "മത്സ്യബന്ധനം"), ("fish_proc", "Fish processing or selling", "മത്സ്യ സംസ്കരണം / വിൽപന"),
        ("manual", "Other manual work", "മറ്റ് കായിക ജോലി"), ("home", "Home duties", "ഗൃഹജോലി"),
        ("student", "Student", "വിദ്യാർത്ഥി"), ("retired", "Retired or not working", "വിരമിച്ചു / ജോലിയില്ല"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("sx", n, en, ml) for n, en, ml in [
        ("none", "None of these", "ഇവയൊന്നുമില്ല"),
        ("diarrhoea", "Loose stools 3 or more times a day", "ദിവസം 3-ൽ കൂടുതൽ തവണ വയറിളക്കം"),
        ("vomiting", "Vomiting", "ഛർദ്ദി"), ("fever", "Fever", "പനി"),
        ("jaundice", "Yellow eyes", "കണ്ണിന് മഞ്ഞനിറം"), ("skin", "Skin rash or boils", "ത്വക്കിൽ ചൊറിച്ചിൽ / കുരുക്കൾ"),
        ("eye", "Red or sticky eyes", "കണ്ണിന് ചുവപ്പ് / പഴുപ്പ്"),
        ("wound", "Wound infected after water contact", "വെള്ളം തട്ടിയ ശേഷം മുറിവ് പഴുത്തു")]]
    ch += [c("sx_care", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("home", "Home remedy", "വീട്ടുവൈദ്യം"),
        ("govt", "PHC or government hospital", "പി.എച്ച്.സി / സർക്കാർ ആശുപത്രി"),
        ("private", "Private clinic or hospital", "സ്വകാര്യ ആശുപത്രി"), ("pharmacy", "Medical shop", "മെഡിക്കൽ ഷോപ്പ്")]]
    ch += [c("cond", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("htn", "High BP", "രക്തസമ്മർദ്ദം (ബി.പി.)"), ("dm", "Diabetes", "പ്രമേഹം"),
        ("ckd", "Kidney disease", "വൃക്കരോഗം"), ("heart", "Heart disease", "ഹൃദ്രോഗം"),
        ("lung", "Asthma or COPD", "ആസ്ത്മ / ശ്വാസകോശ രോഗം"), ("thyroid", "Thyroid", "തൈറോയ്ഡ്"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("tob", n, en, ml) for n, en, ml in [
        ("never", "Never", "ഒരിക്കലുമില്ല"), ("past", "Used in the past", "മുമ്പ് ഉപയോഗിച്ചിരുന്നു"),
        ("current", "Uses now", "ഇപ്പോൾ ഉപയോഗിക്കുന്നു")]]
    settings = {"form_title": "Observatory camp: registration", "form_id": "obs_camp_registration",
                "version": "2026102501", "default_language": "English (en)",
                "instance_name": "concat(${person_id}, ' · ', ${name})"}
    entities = {"list_name": "persons", "label": "concat(${name}, ' (', ${age}, ')')",
                "create_if": "count(instance('persons')/root/item[person_id = ${person_id}]) = 0"}
    return s, ch, settings, entities


# ================================================================ FORM 2

def form_station():
    s = camp_block("station", "Your station today", "ഇന്ന് നിങ്ങളുടെ സ്റ്റേഷൻ")
    s += person_scan_block("Person on the health card", "ഹെൽത്ത് കാർഡിലെ വ്യക്തി")
    s += [
        q("integer", "age", "Age, from the health card", "പ്രായം (ഹെൽത്ത് കാർഡിൽ നിന്ന്)", req="yes",
          con=". >= 0 and . <= 120"),
    ]
    M = "${role}='meas'"
    s += [
        q("begin_group", "meas", "Measurements", "അളവുകൾ", rel=M),
        q("begin_group", "bp", "Blood pressure", "രക്തസമ്മർദ്ദം (ബി.പി.)", rel="${age} >= 18", app="field-list"),
        q("integer", "bp1_sys", "1st reading: systolic (upper)", "ഒന്നാം റീഡിംഗ്: മുകളിലെ (സിസ്റ്റോളിക്)", req="yes",
          con=". >= 60 and . <= 280"),
        q("integer", "bp1_dia", "1st reading: diastolic (lower)", "ഒന്നാം റീഡിംഗ്: താഴത്തെ (ഡയസ്റ്റോളിക്)", req="yes",
          con=". >= 30 and . <= 180 and . < ${bp1_sys}"),
        q("integer", "bp2_sys", "2nd reading after 5 minutes: systolic", "5 മിനിറ്റിനു ശേഷം രണ്ടാം റീഡിംഗ്: മുകളിലെ",
          req="yes", con=". >= 60 and . <= 280"),
        q("integer", "bp2_dia", "2nd reading: diastolic", "രണ്ടാം റീഡിംഗ്: താഴത്തെ", req="yes",
          con=". >= 30 and . <= 180 and . < ${bp2_sys}"),
        q("end_group", "bp"),
        q("calculate", "sbp", calc="round((${bp1_sys} + ${bp2_sys}) div 2, 0)"),
        q("calculate", "dbp", calc="round((${bp1_dia} + ${bp2_dia}) div 2, 0)"),
        q("note", "bp_high", "Average BP ${sbp}/${dbp}: high. Mark the card for the doctor.",
          "ശരാശരി ബി.പി. ${sbp}/${dbp}: കൂടുതലാണ്. ഡോക്ടർക്കായി കാർഡിൽ രേഖപ്പെടുത്തുക.",
          rel="${sbp} >= 140 or ${dbp} >= 90"),
        q("begin_group", "glu", "Blood sugar", "രക്തത്തിലെ പഞ്ചസാര", rel="${age} >= 18", app="field-list"),
        q("integer", "glucose", "Capillary glucose (mg/dL)", "രക്തത്തിലെ പഞ്ചസാര (mg/dL)", req="yes",
          con=". >= 20 and . <= 600"),
        q("select_one glu_type", "glu_type", "Fasting or random?", "വെറും വയറ്റിലാണോ?", req="yes", default="random",
          app="horizontal-compact"),
        q("end_group", "glu"),
        q("begin_group", "anthro", "Height and weight", "ഉയരവും ഭാരവും", app="field-list"),
        q("decimal", "height", "Height (cm)", "ഉയരം (സെ.മീ.)", req="yes", con=". >= 40 and . <= 220"),
        q("decimal", "weight", "Weight (kg)", "ഭാരം (കി.ഗ്രാം)", req="yes", con=". >= 2 and . <= 250"),
        q("end_group", "anthro"),
        q("calculate", "bmi", calc="round(${weight} div ((${height} div 100) * (${height} div 100)), 1)"),
        q("note", "bmi_note", "BMI ${bmi}", "ബി.എം.ഐ ${bmi}", rel="${age} >= 18"),
        q("begin_group", "urine", "Urine dipstick", "മൂത്ര പരിശോധന (ഡിപ്സ്റ്റിക്)", rel="${age} >= 30", app="field-list"),
        q("select_one yes_no", "urine_done", "Urine test done?", "മൂത്ര പരിശോധന നടത്തിയോ?", req="yes",
          app="horizontal-compact"),
        q("select_one dip", "u_protein", "Protein", "പ്രോട്ടീൻ", req="yes", rel="${urine_done}='yes'", app="horizontal-compact"),
        q("select_one dip", "u_glucose", "Glucose", "ഗ്ലൂക്കോസ്", req="yes", rel="${urine_done}='yes'", app="horizontal-compact"),
        q("select_one dip", "u_blood", "Blood", "രക്തം", rel="${urine_done}='yes'", app="horizontal-compact"),
        q("select_one dip", "u_leuk", "Leukocytes", "ല്യൂക്കോസൈറ്റുകൾ", rel="${urine_done}='yes'", app="horizontal-compact"),
        q("select_one posneg", "u_nitrite", "Nitrite", "നൈട്രൈറ്റ്", rel="${urine_done}='yes'", app="horizontal-compact"),
        q("decimal", "u_sg", "Specific gravity", "സ്പെസിഫിക് ഗ്രാവിറ്റി", hen="For example 1.010", hml="ഉദാ: 1.010",
          rel="${urine_done}='yes'", con=". >= 1.000 and . <= 1.040"),
        q("end_group", "urine"),
        q("begin_group", "vitals", "SpO2 and pulse", "ഓക്സിജൻ അളവും നാഡിമിടിപ്പും", rel="${age} >= 60", app="field-list"),
        q("integer", "spo2", "SpO2 (%)", "ഓക്സിജൻ അളവ് (SpO2 %)", con=". >= 50 and . <= 100"),
        q("integer", "pulse", "Pulse (per minute)", "നാഡിമിടിപ്പ് (മിനിറ്റിൽ)", con=". >= 30 and . <= 220"),
        q("end_group", "vitals"),
        q("decimal", "hb", "Haemoglobin (g/dL), if a meter is available", "ഹീമോഗ്ലോബിൻ (g/dL), മീറ്റർ ഉണ്ടെങ്കിൽ",
          con=". >= 3 and . <= 22"),
        q("end_group", "meas"),
    ]
    E = "${role}='eye'"
    s += [
        q("begin_group", "eye", "Eye", "കണ്ണ് പരിശോധന", rel=E),
        q("begin_group", "eye_hx", "Questions", "ചോദ്യങ്ങൾ", app="field-list"),
        q("select_one yes_no", "eye_far", "Difficulty seeing far?", "ദൂരെയുള്ളത് കാണാൻ ബുദ്ധിമുട്ടുണ്ടോ?", req="yes",
          app="horizontal-compact"),
        q("select_one yes_no", "eye_near", "Difficulty seeing near or reading?", "അടുത്തുള്ളത് / വായിക്കാൻ ബുദ്ധിമുട്ടുണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("select_one yes_no", "glasses", "Uses glasses now?", "ഇപ്പോൾ കണ്ണട ഉപയോഗിക്കുന്നുണ്ടോ?", req="yes",
          app="horizontal-compact"),
        q("select_one last_check", "last_eye", "Last eye check", "അവസാനമായി കണ്ണ് പരിശോധിച്ചത്", req="yes"),
        q("select_one yes_no", "dm_known", "Diabetic? (from the card)", "പ്രമേഹമുണ്ടോ? (കാർഡിൽ നിന്ന്)", req="yes",
          app="horizontal-compact"),
        q("integer", "glucose_card", "Blood sugar on the card (mg/dL), if done", "കാർഡിലെ പഞ്ചസാര (mg/dL), ഉണ്ടെങ്കിൽ",
          con=". >= 20 and . <= 600"),
        q("integer", "dm_years", "Years since diabetes was diagnosed", "പ്രമേഹം കണ്ടെത്തിയിട്ട് എത്ര വർഷം?",
          rel="${dm_known}='yes'", con=". >= 0 and . <= 80"),
        q("select_one yes_no", "retina_ever", "Retina ever checked?", "എപ്പോഴെങ്കിലും റെറ്റിന പരിശോധിച്ചിട്ടുണ്ടോ?",
          rel="${dm_known}='yes'", app="horizontal-compact"),
        q("select_one yes_no", "eye_water", "Red or sticky eyes after flood or canal water?",
          "വെള്ളപ്പൊക്ക / കനാൽ വെള്ളം തട്ടിയ ശേഷം കണ്ണിന് ചുവപ്പ് / പഴുപ്പ് ഉണ്ടായോ?", app="horizontal-compact"),
        q("end_group", "eye_hx"),
        q("begin_group", "eye_ex", "Examination", "പരിശോധന", app="field-list"),
        q("select_one va", "va_r", "Distance vision, right eye", "Distance vision, right eye", req="yes"),
        q("select_one va", "va_l", "Distance vision, left eye", "Distance vision, left eye", req="yes"),
        q("select_one yes_no", "pinhole_better", "Improves with pinhole?", "Improves with pinhole?",
          rel="${va_r}!='6_6' or ${va_l}!='6_6'", app="horizontal-compact"),
        q("select_one yes_no", "near_ok", "Near vision normal (N6 at 40 cm)?", "Near vision normal (N6 at 40 cm)?",
          rel="${age} >= 40", app="horizontal-compact"),
        q("select_one yes_no", "readers_given", "Reading glasses given?", "Reading glasses given?",
          rel="${near_ok}='no'", app="horizontal-compact"),
        q("select_one sides", "cataract", "Cataract", "Cataract", rel="${age} >= 40"),
        q("end_group", "eye_ex"),
        q("calculate", "fundus_needed",
          calc="if(${dm_known}='yes' or ${glucose_card} >= 200, 'yes', 'no')"),
        q("begin_group", "fundus", "Retinopathy screening", "Retinopathy screening", rel="${fundus_needed}='yes'",
          app="field-list"),
        q("note", "fundus_why", "Fundus photograph needed: known diabetic or blood sugar 200 or above.",
          "ഫണ്ടസ് ഫോട്ടോ ആവശ്യമാണ്."),
        q("select_one yes_no", "fundus_done", "Fundus photo taken?", "Fundus photo taken?", req="yes",
          app="horizontal-compact"),
        q("text", "fundus_img", "Image ID in the camera app", "Image ID in the camera app",
          hen="Save the image under the person ID in the camera app.", hml="ക്യാമറ ആപ്പിൽ വ്യക്തിയുടെ ഐഡിയിൽ സേവ് ചെയ്യുക.",
          rel="${fundus_done}='yes'"),
        q("select_one dr", "dr_r", "DR grade, right eye", "DR grade, right eye", rel="${fundus_done}='yes'"),
        q("select_one dr", "dr_l", "DR grade, left eye", "DR grade, left eye", rel="${fundus_done}='yes'"),
        q("select_one yes_no", "dme", "Macular oedema suspected?", "Macular oedema suspected?", rel="${fundus_done}='yes'",
          app="horizontal-compact"),
        q("end_group", "fundus"),
        q("decimal", "iop_r", "Eye pressure right (mmHg), if measured", "Eye pressure right (mmHg), if measured",
          rel="${age} >= 40", con=". >= 3 and . <= 70"),
        q("decimal", "iop_l", "Eye pressure left (mmHg), if measured", "Eye pressure left (mmHg), if measured",
          rel="${age} >= 40", con=". >= 3 and . <= 70"),
        q("select_one eye_ref", "eye_referral", "Eye referral", "Eye referral", req="yes"),
        q("end_group", "eye"),
    ]
    D = "${role}='dental'"
    s += [
        q("begin_group", "dental", "Dental", "ദന്ത പരിശോധന", rel=D),
        q("begin_group", "dent_hx", "Questions", "ചോദ്യങ്ങൾ", app="field-list"),
        q("select_multiple oral_sx", "oral_sx", "Any of these now?", "ഇപ്പോൾ ഇവയിൽ ഏതെങ്കിലും ഉണ്ടോ?", req="yes",
          con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("select_one brush", "brush", "How often do you brush?", "എത്ര തവണ പല്ല് തേക്കും?", req="yes", rel="${age} >= 5"),
        q("select_one last_check", "last_dental", "Last dental visit", "അവസാനമായി ദന്തഡോക്ടറെ കണ്ടത്", req="yes"),
        q("select_one tob", "tob_d", "Tobacco, paan or areca use", "പുകയില / മുറുക്ക് / പാക്ക് ഉപയോഗം", req="yes",
          rel="${age} >= 12"),
        q("end_group", "dent_hx"),
        q("begin_group", "dent_ex", "Examination", "പരിശോധന", app="field-list"),
        q("integer", "caries_n", "Number of decayed teeth", "Number of decayed teeth", req="yes", con=". >= 0 and . <= 32"),
        q("integer", "missing_n", "Number of missing teeth", "Number of missing teeth", con=". >= 0 and . <= 32"),
        q("select_one sev4", "gum", "Gum disease", "Gum disease", req="yes"),
        q("select_one hyg", "hygiene", "Oral hygiene", "Oral hygiene", req="yes"),
        q("select_one lesion", "lesion", "Suspicious oral lesion", "Suspicious oral lesion", req="yes",
          rel="${age} >= 30 or ${tob_d}='current'"),
        q("select_one dean", "fluorosis", "Dental fluorosis (Dean's index)", "Dental fluorosis (Dean's index)",
          rel="${age} >= 8 and ${age} <= 15"),
        q("select_one dent_ref", "dent_referral", "Dental referral", "Dental referral", req="yes"),
        q("end_group", "dent_ex"),
        q("end_group", "dental"),
    ]
    G = "${role}='geri'"
    s += [
        q("note", "geri_young", "This screen is for people aged 60 and above.", "ഈ പരിശോധന 60 വയസ്സിനു മുകളിലുള്ളവർക്കു മാത്രം.",
          rel=G + " and ${age} < 60"),
        q("begin_group", "geri", "Geriatric screen (60+)", "വയോജന പരിശോധന (60+)", rel=G + " and ${age} >= 60"),
        q("begin_group", "geri_soc", "Living situation", "ജീവിത സാഹചര്യം", app="field-list"),
        q("select_one yes_no", "alone", "Lives alone?", "ഒറ്റയ്ക്കാണോ താമസം?", req="yes", app="horizontal-compact"),
        q("select_one yes_no", "carer", "Someone available to help when needed?", "ആവശ്യമുള്ളപ്പോൾ സഹായിക്കാൻ ആരെങ്കിലും ഉണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("select_one yes_no", "evac", "Would need help to leave home in a flood?",
          "വെള്ളപ്പൊക്കമുണ്ടായാൽ വീട്ടിൽ നിന്ന് മാറാൻ സഹായം വേണോ?", req="yes", app="horizontal-compact"),
        q("select_one yes_no", "heat_sx", "Felt dizzy or very tired on hot days this year?",
          "ഈ വർഷം ചൂടുള്ള ദിവസങ്ങളിൽ തലകറക്കമോ അമിത ക്ഷീണമോ തോന്നിയോ?", req="yes", app="horizontal-compact"),
        q("integer", "falls", "Falls in the past 12 months", "കഴിഞ്ഞ 12 മാസത്തിൽ എത്ര തവണ വീണു?", req="yes",
          con=". >= 0 and . <= 50"),
        q("end_group", "geri_soc"),
        q("select_one cfs", "cfs", "Clinical Frailty Scale", "Clinical Frailty Scale", req="yes"),
        q("begin_group", "minicog", "Mini-Cog", "Mini-Cog (ഓർമ്മ പരിശോധന)", app="field-list"),
        q("note", "mc_instr",
          "Say three words (banana, sunrise, chair), ask for the clock drawing, then ask for the three words.",
          "മൂന്ന് വാക്കുകൾ പറയുക (വാഴപ്പഴം, സൂര്യോദയം, കസേര), ക്ലോക്ക് വരപ്പിക്കുക, എന്നിട്ട് വാക്കുകൾ ചോദിക്കുക."),
        q("select_one n0to3", "mc_words", "Words recalled", "ഓർമ്മിച്ച വാക്കുകൾ", req="yes", app="horizontal-compact"),
        q("select_one yes_no", "mc_clock", "Clock drawn correctly?", "ക്ലോക്ക് ശരിയായി വരച്ചോ?", req="yes",
          app="horizontal-compact"),
        q("end_group", "minicog"),
        q("calculate", "minicog_score", calc="number(${mc_words}) + if(${mc_clock}='yes', 2, 0)"),
        q("begin_group", "hearing", "Whisper test", "ചെവി പരിശോധന (മന്ത്രണം)", app="field-list"),
        q("select_one passfail", "whisper_r", "Right ear", "വലത് ചെവി", req="yes", app="horizontal-compact"),
        q("select_one passfail", "whisper_l", "Left ear", "ഇടത് ചെവി", req="yes", app="horizontal-compact"),
        q("end_group", "hearing"),
        q("begin_group", "phq", "Mood (PHQ-2)", "മനസ്സിന്റെ അവസ്ഥ (PHQ-2)", app="field-list"),
        q("select_one phq", "phq1", "In the past 2 weeks: little interest or pleasure in doing things?",
          "കഴിഞ്ഞ 2 ആഴ്ചയിൽ: കാര്യങ്ങൾ ചെയ്യുന്നതിൽ താൽപ്പര്യമോ സന്തോഷമോ കുറവായിരുന്നോ?", req="yes"),
        q("select_one phq", "phq2", "Feeling down, depressed or hopeless?", "മനസ്സിന് വിഷമമോ നിരാശയോ തോന്നിയോ?", req="yes"),
        q("end_group", "phq"),
        q("calculate", "phq2_score", calc="number(${phq1}) + number(${phq2})"),
        q("decimal", "grip", "Grip strength, best hand (kg), if a dynamometer is available",
          "കൈപ്പിടി ശക്തി (കി.ഗ്രാം), ഉപകരണം ഉണ്ടെങ്കിൽ", con=". >= 0 and . <= 90"),
        q("note", "geri_flags", "Flag for the doctor: Mini-Cog ${minicog_score}/5, PHQ-2 ${phq2_score}/6.",
          "ഡോക്ടർക്കായി: Mini-Cog ${minicog_score}/5, PHQ-2 ${phq2_score}/6.",
          rel="${minicog_score} < 3 or ${phq2_score} >= 3"),
        q("end_group", "geri"),
    ]
    DR = "${role}='doctor'"
    s += [
        q("begin_group", "doc", "Doctor", "ഡോക്ടർ", rel=DR),
        q("select_multiple dx", "dx", "Diagnosis", "Diagnosis", req="yes"),
        q("text", "dx_other", "Other diagnosis", "Other diagnosis", rel="selected(${dx}, 'other')", req="yes"),
        q("text", "doc_notes", "Notes and treatment", "Notes and treatment", hen=VOICE_EN, hml=VOICE_ML, app="multiline"),
        q("select_one doc_ref", "doc_referral", "Referral", "Referral", req="yes"),
        q("text", "ref_reason", "Reason for referral", "Reason for referral", rel="${doc_referral}!='none'", req="yes"),
        q("select_one yes_no", "followup", "Follow-up needed by ASHA?", "Follow-up needed by ASHA?", req="yes",
          app="horizontal-compact"),
        q("end_group", "doc"),
        q("note", "st_done", "Saved for ${person_id}. Hand the card back.", "${person_id} സേവ് ചെയ്തു. കാർഡ് തിരികെ നൽകുക."),
    ]

    ch = common_choices()
    ch += [c("station", n, en, ml) for n, en, ml in [
        ("meas", "Measurements", "അളവുകൾ"), ("eye", "Eye", "കണ്ണ്"), ("dental", "Dental", "ദന്തം"),
        ("geri", "Geriatric screen (MSW)", "വയോജന പരിശോധന (MSW)"), ("doctor", "Doctor", "ഡോക്ടർ")]]
    ch += [c("glu_type", n, en, ml) for n, en, ml in [("random", "Random", "റാൻഡം"), ("fasting", "Fasting", "വെറും വയറ്റിൽ")]]
    ch += [c("dip", n, en) for n, en in [("neg", "Neg"), ("trace", "Trace"), ("p1", "1+"), ("p2", "2+"), ("p3", "3+")]]
    ch += [c("posneg", n, en) for n, en in [("neg", "Negative"), ("pos", "Positive")]]
    ch += [c("last_check", n, en, ml) for n, en, ml in [
        ("never", "Never", "ഒരിക്കലുമില്ല"), ("lt2y", "Within 2 years", "2 വർഷത്തിനുള്ളിൽ"),
        ("gt2y", "Over 2 years ago", "2 വർഷത്തിലേറെ മുമ്പ്")]]
    ch += [c("va", n, en) for n, en in [
        ("6_6", "6/6"), ("6_9", "6/9"), ("6_12", "6/12"), ("6_18", "6/18"), ("6_24", "6/24"), ("6_36", "6/36"),
        ("6_60", "6/60"), ("lt6_60", "Worse than 6/60"), ("cf", "Counting fingers"), ("hm", "Hand movements"),
        ("pl", "Perceives light"), ("npl", "No light perception")]]
    ch += [c("sides", n, en) for n, en in [("none", "None"), ("right", "Right"), ("left", "Left"), ("both", "Both")]]
    ch += [c("dr", n, en) for n, en in [
        ("none", "No DR"), ("mild", "Mild NPDR"), ("moderate", "Moderate NPDR"), ("severe", "Severe NPDR"),
        ("pdr", "Proliferative DR"), ("ungradable", "Ungradable")]]
    ch += [c("eye_ref", n, en) for n, en in [
        ("none", "None"), ("glasses", "Spectacles"), ("cataract", "Cataract surgery"), ("retina", "Retina specialist"),
        ("glaucoma", "Glaucoma evaluation"), ("other", "Other")]]
    ch += [c("oral_sx", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("pain", "Tooth pain", "പല്ലുവേദന"), ("bleed", "Bleeding gums", "മോണയിൽ നിന്ന് രക്തം"),
        ("ulcer", "Mouth ulcers", "വായിൽ പുണ്ണ്")]]
    ch += [c("brush", n, en, ml) for n, en, ml in [
        ("lt1", "Less than daily", "ദിവസവും ഇല്ല"), ("once", "Once a day", "ദിവസം ഒരു തവണ"),
        ("twice", "Twice a day", "ദിവസം രണ്ട് തവണ")]]
    ch += [c("tob", n, en, ml) for n, en, ml in [
        ("never", "Never", "ഒരിക്കലുമില്ല"), ("past", "Used in the past", "മുമ്പ് ഉപയോഗിച്ചിരുന്നു"),
        ("current", "Uses now", "ഇപ്പോൾ ഉപയോഗിക്കുന്നു")]]
    ch += [c("sev4", n, en) for n, en in [("none", "None"), ("mild", "Mild"), ("moderate", "Moderate"), ("severe", "Severe")]]
    ch += [c("hyg", n, en) for n, en in [("good", "Good"), ("fair", "Fair"), ("poor", "Poor")]]
    ch += [c("lesion", n, en) for n, en in [
        ("none", "None"), ("leuko", "White patch (leukoplakia)"), ("erythro", "Red patch (erythroplakia)"),
        ("ulcer", "Ulcer over 2 weeks"), ("osmf", "Submucous fibrosis"), ("other", "Other")]]
    ch += [c("dean", n, en) for n, en in [
        ("normal", "Normal"), ("questionable", "Questionable"), ("very_mild", "Very mild"), ("mild", "Mild"),
        ("moderate", "Moderate"), ("severe", "Severe")]]
    ch += [c("dent_ref", n, en) for n, en in [
        ("none", "None"), ("treatment", "Dental treatment"), ("oral_med", "Oral medicine or biopsy"), ("other", "Other")]]
    ch += [c("cfs", str(i), en) for i, en in [
        (1, "1 Very fit"), (2, "2 Fit"), (3, "3 Managing well"), (4, "4 Living with very mild frailty"),
        (5, "5 Living with mild frailty"), (6, "6 Living with moderate frailty"), (7, "7 Living with severe frailty"),
        (8, "8 Living with very severe frailty"), (9, "9 Terminally ill")]]
    ch += [c("n0to3", str(i), str(i)) for i in range(4)]
    ch += [c("passfail", n, en, ml) for n, en, ml in [("pass", "Heard", "കേട്ടു"), ("fail", "Not heard", "കേട്ടില്ല")]]
    ch += [c("phq", str(i), en, ml) for i, en, ml in [
        (0, "Not at all", "ഇല്ല"), (1, "Several days", "ചില ദിവസങ്ങൾ"),
        (2, "More than half the days", "പകുതിയിലധികം ദിവസങ്ങൾ"), (3, "Nearly every day", "മിക്കവാറും എല്ലാ ദിവസവും")]]
    ch += [c("dx", n, en) for n, en in [
        ("htn_new", "Hypertension, new"), ("htn_unc", "Hypertension, uncontrolled"), ("dm_new", "Diabetes, new"),
        ("dm_unc", "Diabetes, uncontrolled"), ("diarrhoea", "Acute diarrhoea"), ("skin", "Skin infection"),
        ("uri", "Respiratory infection"), ("anaemia", "Anaemia"), ("msk", "Musculoskeletal pain"),
        ("ckd_susp", "Kidney disease suspected"), ("cog", "Cognitive impairment suspected"),
        ("depr", "Depression screen positive"), ("frail", "Frailty"), ("lepto_susp", "Leptospirosis suspected"),
        ("well", "No illness found"), ("other", "Other")]]
    ch += [c("doc_ref", n, en) for n, en in [
        ("none", "None"), ("phc", "PHC / FHC"), ("taluk", "Taluk or district hospital"),
        ("mch", "Medical college"), ("specialist", "Specialist clinic")]]
    settings = {"form_title": "Observatory camp: station record", "form_id": "obs_camp_station",
                "version": "2026102401", "default_language": "English (en)",
                "instance_name": "concat(${person_id}, ' · ', ${role})"}
    return s, ch, settings


# ---------------------------------------------------------------- writer

HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="12262E")
BODY_FONT = Font(name="Arial")
GROUP_FILL = PatternFill("solid", fgColor="E7F3F2")


def write(path, survey, choices, settings, entities=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(SURVEY_COLS)
    for r in survey:
        ws.append(r)
    wc = wb.create_sheet("choices")
    wc.append(CHOICE_COLS)
    for r in choices:
        wc.append(r)
    wst = wb.create_sheet("settings")
    keys = list(settings)
    wst.append(keys)
    wst.append([settings[k] for k in keys])
    sheets = [(ws, [22, 18, 48, 48, 30, 30, 9, 30, 30, 26, 26, 34, 24, 18, 12, 24, 14]),
              (wc, [14, 14, 36, 36, 12]), (wst, [34] * len(keys))]
    if entities:
        we = wb.create_sheet("entities")
        ekeys = list(entities)
        we.append(ekeys)
        we.append([entities[k] for k in ekeys])
        sheets.append((we, [16, 40, 60]))
    for sheet, widths in sheets:
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = w
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = HEAD_FONT
            cell.fill = HEAD_FILL
        sheet.freeze_panes = "C2" if sheet is ws else "A2"
    for row in ws.iter_rows(min_row=2):
        if row[0].value in ("begin_group", "end_group"):
            for cell in row:
                cell.fill = GROUP_FILL
    wb.save(path)


def hh_scan_block(rel="", save=False):
    return [
        q("begin_group", "hid", "Household ID", "വീടിന്റെ ഐഡി", rel=rel, app="field-list"),
        q("select_one id_method", "hh_method", "Household ID: scan or type?", "വീടിന്റെ ഐഡി: സ്കാൻ ചെയ്യണോ ടൈപ്പ് ചെയ്യണോ?",
          req="yes", default="scan"),
        q("barcode", "hh_scan", "Scan the household sticker", "വീടിന്റെ സ്റ്റിക്കർ സ്കാൻ ചെയ്യുക", req="yes",
          rel="${hh_method}='scan'", con=HH_RE, cmsg_en="This is not a household sticker (format CVB-014).",
          cmsg_ml="ഇത് വീടിന്റെ സ്റ്റിക്കർ അല്ല (ഉദാ: CVB-014)."),
        q("text", "hh_type", "Type the household ID", "വീടിന്റെ ഐഡി ടൈപ്പ് ചെയ്യുക", hen="Format: CVB-014",
          hml="ഉദാ: CVB-014", req="yes", rel="${hh_method}='type'", con=HH_RE,
          cmsg_en="Use the format CVB-014.", cmsg_ml="CVB-014 എന്ന രൂപത്തിൽ നൽകുക."),
        q("end_group", "hid"),
        q("calculate", "hh_id", calc="if(${hh_method}='scan', ${hh_scan}, ${hh_type})", save="hh_id" if save else ""),
    ]


FIELD_ROLES = [("asha", "ASHA", "ആശാ പ്രവർത്തക"), ("volunteer", "Volunteer", "വോളണ്ടിയർ"),
               ("oh_vol", "One Health volunteer", "വൺ ഹെൽത്ത് വോളണ്ടിയർ"), ("msw", "MSW", "എം.എസ്.ഡബ്ല്യു"),
               ("nio", "NIO team", "എൻ.ഐ.ഒ. ടീം"), ("coord", "Coordinator", "കോർഡിനേറ്റർ")]


# ================================================================ FORM 3

def form_home_visit():
    s = camp_block("field_role")
    s += hh_scan_block()
    s += [
        q("select_multiple visit_reason", "reason", "Reason for this visit", "ഈ സന്ദർശനത്തിന്റെ കാരണം", req="yes"),
        q("begin_group", "illness", "Illness in this household, past 2 weeks", "ഈ വീട്ടിൽ കഴിഞ്ഞ 2 ആഴ്ചയിലെ അസുഖങ്ങൾ",
          app="field-list"),
        q("select_one yes_no", "anyone_ill", "Was anyone in the house ill in the past 2 weeks?",
          "കഴിഞ്ഞ 2 ആഴ്ചയിൽ വീട്ടിൽ ആർക്കെങ്കിലും അസുഖമുണ്ടായോ?", req="yes", app="horizontal-compact"),
        q("integer", "n_diarrhoea", "People with loose stools 3+ times a day", "ദിവസം 3-ൽ കൂടുതൽ തവണ വയറിളക്കമുള്ളവർ",
          req="yes", rel="${anyone_ill}='yes'", con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_diarrhoea_u5", "Of these, children under 5", "ഇതിൽ 5 വയസ്സിൽ താഴെയുള്ള കുട്ടികൾ",
          rel="${anyone_ill}='yes' and ${n_diarrhoea} > 0", con=". >= 0 and . <= ${n_diarrhoea}", default="0"),
        q("integer", "n_fever", "People with fever", "പനിയുള്ളവർ", req="yes", rel="${anyone_ill}='yes'",
          con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_jaundice", "People with yellow eyes", "കണ്ണിന് മഞ്ഞനിറമുള്ളവർ", rel="${anyone_ill}='yes'",
          con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_skin", "Skin rash or boils", "ത്വക്കിൽ ചൊറിച്ചിൽ / കുരുക്കൾ", rel="${anyone_ill}='yes'",
          con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_eye", "Red or sticky eyes", "കണ്ണിന് ചുവപ്പ് / പഴുപ്പ്", rel="${anyone_ill}='yes'",
          con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_wound", "Wound infected after water contact", "വെള്ളം തട്ടിയ ശേഷം മുറിവ് പഴുത്തവർ",
          rel="${anyone_ill}='yes'", con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_heat", "Elders 60+ dizzy or exhausted in the heat", "ചൂടിൽ തലകറക്കം / ക്ഷീണം തോന്നിയ 60+ പ്രായക്കാർ",
          rel="${anyone_ill}='yes'", con=". >= 0 and . <= 30", default="0"),
        q("integer", "n_hosp", "Admitted to hospital", "ആശുപത്രിയിൽ പ്രവേശിപ്പിച്ചവർ", rel="${anyone_ill}='yes'",
          con=". >= 0 and . <= 30", default="0"),
        q("end_group", "illness"),
        q("begin_group", "refgrp", "Referral follow-up", "റഫറൽ തുടർനടപടി", rel="selected(${reason}, 'referral')",
          app="field-list"),
        q("text", "ref_person", "Person ID on the health card", "ഹെൽത്ത് കാർഡിലെ വ്യക്തിയുടെ ഐഡി", hen="Format: OBS-P-000123",
          hml="ഉദാ: OBS-P-000123", req="yes", con=PERSON_RE, cmsg_en="Use the format OBS-P-000123.",
          cmsg_ml="OBS-P-000123 എന്ന രൂപത്തിൽ നൽകുക."),
        q("select_one ref_status", "ref_status", "Did they go to the referral?", "റഫർ ചെയ്ത സ്ഥലത്ത് പോയോ?", req="yes"),
        q("select_multiple barrier", "ref_barrier", "What stopped them?", "എന്താണ് തടസ്സം?",
          rel="${ref_status}='not_yet' or ${ref_status}='refused'"),
        q("end_group", "refgrp"),
        q("begin_group", "eldgrp", "Elder check (tide, flood or heat)", "വയോജന പരിശോധന (വേലിയേറ്റം, വെള്ളപ്പൊക്കം, ചൂട്)",
          rel="selected(${reason}, 'elder')", app="field-list"),
        q("integer", "eld_checked", "Elders 60+ checked in this house", "ഈ വീട്ടിൽ പരിശോധിച്ച 60+ പ്രായക്കാർ", req="yes",
          con=". >= 1 and . <= 10"),
        q("select_one yes_no", "eld_meds", "Medicines for at least 5 days at home?", "കുറഞ്ഞത് 5 ദിവസത്തേക്കുള്ള മരുന്ന് വീട്ടിലുണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("select_one yes_no", "eld_phone", "A charged phone with the ASHA's number?", "ആശാ പ്രവർത്തകയുടെ നമ്പറുള്ള ചാർജുള്ള ഫോൺ ഉണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("select_one yes_no", "eld_evac", "Needs help to move out if water rises?", "വെള്ളം കയറിയാൽ മാറാൻ സഹായം വേണോ?",
          req="yes", app="horizontal-compact"),
        q("select_one yes_no", "eld_water", "Drinking enough water, has shade or a fan?", "ആവശ്യത്തിന് വെള്ളം കുടിക്കുന്നുണ്ടോ, തണലോ ഫാനോ ഉണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("end_group", "eldgrp"),
        q("begin_group", "advgrp", "Advice given", "നൽകിയ നിർദ്ദേശങ്ങൾ", app="field-list"),
        q("select_multiple advice", "advice", "Advice or results given today", "ഇന്ന് നൽകിയ നിർദ്ദേശങ്ങൾ / ഫലങ്ങൾ", req="yes",
          con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("integer", "chlorine_given", "Chlorine tablets given", "നൽകിയ ക്ലോറിൻ ഗുളികകൾ", rel="selected(${advice}, 'chlorinate')",
          con=". >= 0 and . <= 500"),
        q("end_group", "advgrp"),
        q("text", "notes", "Notes", "കുറിപ്പുകൾ", hen=VOICE_EN, hml=VOICE_ML, app="multiline"),
    ]
    ch = common_choices()
    ch += [c("field_role", n, en, ml) for n, en, ml in FIELD_ROLES]
    ch += [c("visit_reason", n, en, ml) for n, en, ml in [
        ("routine", "Routine visit", "സാധാരണ സന്ദർശനം"), ("referral", "Referral follow-up", "റഫറൽ തുടർനടപടി"),
        ("result", "Give water result or advice", "വെള്ള പരിശോധനാ ഫലം / നിർദ്ദേശം നൽകൽ"),
        ("elder", "Elder check before tide, flood or heat", "വേലിയേറ്റം / വെള്ളപ്പൊക്കം / ചൂടിന് മുമ്പ് വയോജന പരിശോധന"),
        ("outbreak", "Outbreak or cluster check", "രോഗവ്യാപന പരിശോധന")]]
    ch += [c("ref_status", n, en, ml) for n, en, ml in [
        ("went", "Went", "പോയി"), ("not_yet", "Not yet", "ഇതുവരെ പോയില്ല"), ("refused", "Refused", "പോകാൻ വിസമ്മതിച്ചു"),
        ("moved", "Moved away or not found", "താമസം മാറി / കണ്ടില്ല")]]
    ch += [c("barrier", n, en, ml) for n, en, ml in [
        ("cost", "Cost", "ചെലവ്"), ("transport", "Transport", "യാത്രാസൗകര്യം"), ("no_escort", "No one to go with", "കൂടെ പോകാൻ ആളില്ല"),
        ("work", "Work or wage loss", "ജോലി / കൂലി നഷ്ടം"), ("fear", "Fear or not needed", "ഭയം / ആവശ്യമില്ലെന്ന് തോന്നി"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("advice", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("result", "Water result explained", "വെള്ള പരിശോധനാ ഫലം വിശദീകരിച്ചു"),
        ("boil", "Boil drinking water", "കുടിവെള്ളം തിളപ്പിക്കുക"), ("chlorinate", "Chlorinate well or water", "കിണർ / വെള്ളം ക്ലോറിനേറ്റ് ചെയ്യുക"),
        ("tank", "Clean the tank or sump", "ടാങ്ക് / സംപ് വൃത്തിയാക്കുക"), ("heat", "Heat advice", "ചൂട് സംബന്ധിച്ച നിർദ്ദേശം"),
        ("flood", "Flood or tide preparation", "വെള്ളപ്പൊക്ക / വേലിയേറ്റ തയ്യാറെടുപ്പ്"),
        ("ors", "ORS and when to see a doctor", "ഒ.ആർ.എസ്, എപ്പോൾ ഡോക്ടറെ കാണണം")]]
    settings = {"form_title": "Observatory: household visit", "form_id": "obs_household_visit",
                "version": "2026102401", "default_language": "English (en)",
                "instance_name": "concat(${hh_id}, ' · ', ${today})"}
    return s, ch, settings


# ================================================================ FORM 4

def form_water_sample():
    s = camp_block("field_role")
    s += [
        q("select_one point_type", "point_type", "Where is this sample from?", "ഈ സാമ്പിൾ എവിടെ നിന്ന്?", req="yes",
          default="${last-saved#point_type}"),
    ]
    s += hh_scan_block(rel="${point_type}='household'")
    s += [
        q("begin_group", "sample", "Sample", "സാമ്പിൾ", app="field-list"),
        q("barcode", "sample_id", "Scan the bottle label", "കുപ്പിയിലെ ലേബൽ സ്കാൻ ചെയ്യുക", req="yes",
          con="regex(., '^WS-[0-9]{6}$')", cmsg_en="This is not a sample bottle label (format WS-000123).",
          cmsg_ml="ഇത് സാമ്പിൾ കുപ്പിയുടെ ലേബൽ അല്ല (ഉദാ: WS-000123)."),
        q("select_one sample_src", "source", "Water source", "വെള്ളത്തിന്റെ ഉറവിടം", req="yes"),
        q("select_one sample_point", "taken_at", "Taken from", "എവിടെ നിന്ന് എടുത്തു", req="yes"),
        q("geopoint", "gps", "Location of the source", "ഉറവിടത്തിന്റെ സ്ഥാനം", req="yes",
          hen="Stand next to the well or tap.", hml="കിണറിനോ ടാപ്പിനോ അടുത്ത് നിൽക്കുക."),
        q("end_group", "sample"),
        q("calculate", "sampled_at", calc="now()"),
        q("begin_group", "field", "Field readings", "സ്ഥലത്തെ അളവുകൾ", app="field-list"),
        q("integer", "tds", "TDS (mg/L)", "TDS (mg/L)", con=". >= 0 and . <= 50000"),
        q("decimal", "water_temp", "Water temperature (°C)", "വെള്ളത്തിന്റെ താപനില (°C)", con=". >= 10 and . <= 45"),
        q("select_one clarity", "clarity", "How does it look?", "വെള്ളം എങ്ങനെ കാണപ്പെടുന്നു?"),
        q("select_one smell", "smell", "Smell", "മണം"),
        q("integer", "fu_colour", "Forel-Ule colour (1–21), if a disk is used", "ഫോറൽ-ഊലെ നിറം (1–21), ഡിസ്ക് ഉണ്ടെങ്കിൽ",
          con=". >= 1 and . <= 21"),
        q("select_one yes_no", "rain_48h", "Heavy rain in the past 48 hours?", "കഴിഞ്ഞ 48 മണിക്കൂറിൽ കനത്ത മഴ പെയ്തോ?", req="yes",
          app="horizontal-compact"),
        q("select_one yes_no", "flooded_7d", "Source flooded in the past 7 days?", "കഴിഞ്ഞ 7 ദിവസത്തിൽ ഉറവിടത്തിൽ വെള്ളം കയറിയോ?",
          req="yes", app="horizontal-compact"),
        q("end_group", "field"),
        q("begin_group", "aq", "Aquadip©", "Aquadip©", app="field-list"),
        q("select_one aq_result", "aquadip", "Aquadip© result, if tested here", "Aquadip© ഫലം, ഇവിടെ പരിശോധിച്ചെങ്കിൽ", req="yes"),
        q("image", "aq_photo", "Photo of the Aquadip© result", "Aquadip© ഫലത്തിന്റെ ഫോട്ടോ",
          rel="${aquadip}='pos' or ${aquadip}='neg'", params="max-pixels=1024"),
        q("select_one yes_no", "to_lab", "Sent to the NIO lab?", "എൻ.ഐ.ഒ. ലാബിലേക്ക് അയച്ചോ?", req="yes",
          app="horizontal-compact"),
        q("end_group", "aq"),
        q("text", "notes", "Notes", "കുറിപ്പുകൾ", hen=VOICE_EN, hml=VOICE_ML, app="multiline"),
    ]
    ch = common_choices()
    ch += [c("field_role", n, en, ml) for n, en, ml in FIELD_ROLES]
    ch += [c("point_type", n, en, ml) for n, en, ml in [
        ("household", "A household", "ഒരു വീട്"), ("public", "Public well, tap or tank", "പൊതു കിണർ / ടാപ്പ് / ടാങ്ക്"),
        ("canal", "Canal, backwater or sea", "കനാൽ / കായൽ / കടൽ"), ("tanker", "Tanker fill point", "ടാങ്കർ നിറയ്ക്കുന്ന സ്ഥലം")]]
    ch += [c("sample_src", n, en, ml) for n, en, ml in [
        ("own_well", "Own well", "സ്വന്തം കിണർ"), ("shared_well", "Shared or public well", "പൊതു / പങ്കിട്ട കിണർ"),
        ("tap", "Tap (KWA / piped)", "പൈപ്പ് വെള്ളം (KWA)"), ("tanker", "Tanker", "ടാങ്കർ വെള്ളം"),
        ("borewell", "Borewell", "കുഴൽക്കിണർ"), ("sump", "Sump", "സംപ്"), ("overhead", "Overhead tank", "മുകളിലെ ടാങ്ക്"),
        ("canal", "Canal or backwater", "കനാൽ / കായൽ"), ("sea", "Sea", "കടൽ"), ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("sample_point", n, en, ml) for n, en, ml in [
        ("source", "Directly from the source", "ഉറവിടത്തിൽ നിന്ന് നേരിട്ട്"),
        ("stored", "Stored drinking water in the house", "വീട്ടിൽ സൂക്ഷിച്ച കുടിവെള്ളം")]]
    ch += [c("clarity", n, en, ml) for n, en, ml in [
        ("clear", "Clear", "തെളിഞ്ഞത്"), ("slight", "Slightly cloudy", "അൽപം കലങ്ങിയത്"), ("turbid", "Muddy or cloudy", "കലങ്ങിയത്"),
        ("coloured", "Coloured", "നിറമുള്ളത്")]]
    ch += [c("smell", n, en, ml) for n, en, ml in [
        ("none", "None", "ഇല്ല"), ("sewage", "Sewage", "മലിനജല ഗന്ധം"), ("rotten_egg", "Rotten egg", "ചീമുട്ട ഗന്ധം"),
        ("chlorine", "Chlorine", "ക്ലോറിൻ"), ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("aq_result", n, en, ml) for n, en, ml in [
        ("pos", "Positive (contaminated)", "പോസിറ്റീവ് (മലിനം)"), ("neg", "Negative", "നെഗറ്റീവ്"),
        ("not_done", "Not tested here", "ഇവിടെ പരിശോധിച്ചില്ല")]]
    settings = {"form_title": "Observatory: water sample", "form_id": "obs_water_sample",
                "version": "2026102401", "default_language": "English (en)",
                "instance_name": "concat(${sample_id}, ' · ', ${camp_site})"}
    return s, ch, settings


# ================================================================ REPORT FORMS
# Any form whose ID starts with obs_report_ is a "report form". The Observatory reads the field
# names below from every report form, so a new report form made by copying one of these keeps working.

REPORT_ROLES = FIELD_ROLES + [("ward", "Ward member", "വാർഡ് മെമ്പർ")]
DANGER_NONE = "not(selected(${dehyd}, 'none'))"


def report_start(form_title_en, form_title_ml):
    """Reporter, household (scanned from the door sticker) and the site derived from it."""
    return [
        q("start", "start"), q("end", "end"), q("today", "today"), q("deviceid", "deviceid"),
        q("csv-external", "persons"),
        q("note", "intro", form_title_en + ". Scan the sticker on the house door to begin. "
          "(Malayalam text is a machine-assisted translation for accessibility; English is the reference.)",
          form_title_ml + ". തുടങ്ങാൻ വീടിന്റെ വാതിലിലെ സ്റ്റിക്കർ സ്കാൻ ചെയ്യുക. "
          "(മലയാളം യന്ത്രസഹായത്തോടെയുള്ള വിവർത്തനമാണ്; ഇംഗ്ലീഷ് പാഠമാണ് ആധികാരികം.)"),
        q("select_one report_role", "role", "You are reporting as", "നിങ്ങൾ റിപ്പോർട്ട് ചെയ്യുന്നത്", req="yes",
          default="${last-saved#role}"),
    ] + hh_scan_block(save=True) + [
        q("calculate", "camp_site", calc="substring-before(${hh_id}, '-')"),
        q("calculate", "n_members", calc="count(instance('persons')/root/item[hh_id = ${hh_id}])"),
    ]


def report_person():
    """Consent, then pick the person from the household list or add them. Without consent: age and sex only."""
    ON_LIST = "${in_list_eff}='yes'"
    NEW = "${consent_named}='yes' and ${in_list_eff}='no'"
    return [
        q("select_one yes_no", "consent_named",
          "Does the family agree to share this report, with the person's name, with the health team (ASHA and doctor)?",
          "ഈ റിപ്പോർട്ട് പേര് സഹിതം ആരോഗ്യ സംഘവുമായി (ആശാ പ്രവർത്തക, ഡോക്ടർ) പങ്കുവയ്ക്കാൻ കുടുംബം സമ്മതിക്കുന്നുണ്ടോ?",
          req="yes", app="horizontal-compact"),
        q("note", "anon_note", "That is fine. The report will be kept without a name.",
          "കുഴപ്പമില്ല. പേരില്ലാതെ റിപ്പോർട്ട് സൂക്ഷിക്കും.", rel="${consent_named}='no'"),
        q("begin_group", "who", "Who is ill?", "ആർക്കാണ് അസുഖം?", rel="${consent_named}='yes'", app="field-list"),
        q("select_one in_list", "in_list", "Is the person on this household's list?",
          "ഈ വ്യക്തി വീട്ടിലെ പട്ടികയിൽ ഉണ്ടോ?", req="yes", rel="${n_members} > 0"),
        q("select_one_from_file persons.csv", "person", "Choose the person", "വ്യക്തിയെ തിരഞ്ഞെടുക്കുക", req="yes",
          rel="${n_members} > 0 and ${in_list}='yes'", cf="hh_id = ${hh_id}"),
        q("note", "none_listed", "No one from this house is on the list yet. Add the person below.",
          "ഈ വീട്ടിൽ നിന്ന് ആരും പട്ടികയിൽ ഇല്ല. താഴെ ചേർക്കുക.", rel="${n_members} = 0"),
        q("end_group", "who"),
        q("calculate", "in_list_eff", calc="if(${n_members} = 0, 'no', ${in_list})"),
        q("begin_group", "newp", "Add a person to this household", "ഈ വീട്ടിലേക്ക് ഒരാളെ ചേർക്കുക", rel=NEW,
          app="field-list"),
        q("text", "new_name", "Name", "പേര്", hen="Manglish or Malayalam keyboard both work.",
          hml="മംഗ്ലീഷ് അല്ലെങ്കിൽ മലയാളം കീബോർഡ് ഉപയോഗിക്കാം.", req="yes", save="full_name"),
        q("integer", "new_age", "Age (years)", "പ്രായം (വർഷം)", req="yes", con=". >= 0 and . <= 120", save="age"),
        q("select_one sex", "new_sex", "Sex", "ലിംഗം", req="yes", app="horizontal-compact", save="sex"),
        q("barcode", "new_card", "If they have a person ID card from a camp, scan it (optional)",
          "ക്യാമ്പിലെ ഐഡി കാർഡ് ഉണ്ടെങ്കിൽ സ്കാൻ ചെയ്യുക (ഐച്ഛികം)", con=PERSON_RE,
          cmsg_en="This is not a person ID sticker (format OBS-P-000123).",
          cmsg_ml="ഇത് വ്യക്തിയുടെ ഐഡി സ്റ്റിക്കർ അല്ല (ഉദാ: OBS-P-000123)."),
        q("end_group", "newp"),
        q("calculate", "new_pid", calc="if(string-length(${new_card}) > 0, ${new_card}, "
          "once(concat('OBS-N-', substr(uuid(), 0, 6))))", save="person_id"),
        q("begin_group", "anon", "About the person", "വ്യക്തിയെ കുറിച്ച്", rel="${consent_named}='no'", app="field-list"),
        q("integer", "anon_age", "Age (years)", "പ്രായം (വർഷം)", req="yes", con=". >= 0 and . <= 120"),
        q("select_one sex", "anon_sex", "Sex", "ലിംഗം", req="yes", app="horizontal-compact"),
        q("end_group", "anon"),
        q("calculate", "person_id", calc=f"if(${{consent_named}}!='yes', '', if({ON_LIST}, "
          "instance('persons')/root/item[name = ${person}]/person_id, ${new_pid}))"),
        q("calculate", "p_age", calc=f"if(${{consent_named}}!='yes', ${{anon_age}}, if({ON_LIST}, "
          "instance('persons')/root/item[name = ${person}]/age, ${new_age}))"),
        q("calculate", "p_sex", calc=f"if(${{consent_named}}!='yes', ${{anon_sex}}, if({ON_LIST}, "
          "instance('persons')/root/item[name = ${person}]/sex, ${new_sex}))"),
    ]


def report_impact():
    """Days lost: the same questions on every illness report, so workdays lost add up across forms."""
    WORKS = "selected(${activity}, 'daily_wage') or selected(${activity}, 'fishing') or selected(${activity}, 'salaried') " \
            "or selected(${activity}, 'self_emp') or selected(${activity}, 'home')"
    return [
        q("begin_group", "impact", "Effect on daily life", "ദൈനംദിന ജീവിതത്തെ ബാധിച്ചത്", app="field-list"),
        q("date", "onset", "When did it start?", "എപ്പോഴാണ് തുടങ്ങിയത്?", req="yes",
          con=". <= today() and . >= today() - 60", cmsg_en="Within the past 60 days, and not in the future.",
          cmsg_ml="കഴിഞ്ഞ 60 ദിവസത്തിനുള്ളിൽ ആയിരിക്കണം."),
        q("select_one yes_no", "still_ill", "Still ill today?", "ഇന്നും അസുഖമുണ്ടോ?", req="yes", app="horizontal-compact"),
        q("integer", "days_ill", "Days ill so far", "ഇതുവരെ എത്ര ദിവസം അസുഖം?", req="yes", con=". >= 0 and . <= 90"),
        q("select_one activity", "activity", "What does this person usually do?", "ഈ വ്യക്തി സാധാരണ എന്താണ് ചെയ്യുന്നത്?",
          req="yes"),
        q("integer", "work_days", "Days unable to work (including home duties)", "ജോലി ചെയ്യാനാകാത്ത ദിവസങ്ങൾ (വീട്ടുജോലി ഉൾപ്പെടെ)",
          req="yes", rel=WORKS, con=". >= 0 and . <= ${days_ill}",
          cmsg_en="Cannot be more than the days ill.", cmsg_ml="അസുഖ ദിവസങ്ങളേക്കാൾ കൂടരുത്."),
        q("integer", "school_days", "School or college days missed", "നഷ്ടപ്പെട്ട സ്കൂൾ / കോളേജ് ദിവസങ്ങൾ", req="yes",
          rel="selected(${activity}, 'student')", con=". >= 0 and . <= ${days_ill}",
          cmsg_en="Cannot be more than the days ill.", cmsg_ml="അസുഖ ദിവസങ്ങളേക്കാൾ കൂടരുത്."),
        q("integer", "carer_days", "Days a family member stayed away from work to look after them",
          "ഇവരെ നോക്കാൻ വീട്ടിലെ ആൾ ജോലിക്ക് പോകാതിരുന്ന ദിവസങ്ങൾ", req="yes", default="0", con=". >= 0 and . <= 90"),
        q("select_one care", "care", "Treatment taken", "എടുത്ത ചികിത്സ", req="yes"),
        q("integer", "cost_rs", "Money spent on treatment and travel (₹), if known",
          "ചികിത്സയ്ക്കും യാത്രയ്ക്കും ചെലവായ തുക (₹), അറിയാമെങ്കിൽ", con=". >= 0 and . <= 1000000"),
        q("end_group", "impact"),
        q("text", "notes", "Notes", "കുറിപ്പുകൾ", hen=VOICE_EN, hml=VOICE_ML, app="multiline"),
    ]


def report_choices():
    ch = common_choices()
    ch += [c("report_role", n, en, ml) for n, en, ml in REPORT_ROLES]
    ch += [c("in_list", "yes", "Yes, choose from the list", "ഉണ്ട്, പട്ടികയിൽ നിന്ന് തിരഞ്ഞെടുക്കാം"),
           c("in_list", "no", "No, add a new person", "ഇല്ല, പുതിയ ആളെ ചേർക്കാം")]
    ch += [c("sex", n, en, ml) for n, en, ml in [("m", "Male", "പുരുഷൻ"), ("f", "Female", "സ്ത്രീ"), ("o", "Other", "മറ്റുള്ളവ")]]
    ch += [c("activity", n, en, ml) for n, en, ml in [
        ("daily_wage", "Daily wage work", "ദിവസക്കൂലി ജോലി"), ("fishing", "Fishing or fish work", "മത്സ്യബന്ധനം / മത്സ്യ ജോലി"),
        ("salaried", "Salaried job", "ശമ്പള ജോലി"), ("self_emp", "Own business or farming", "സ്വന്തം സംരംഭം / കൃഷി"),
        ("home", "Home duties", "വീട്ടുജോലി"), ("student", "Student", "വിദ്യാർത്ഥി"),
        ("child", "Young child, not in school", "സ്കൂളിൽ പോകാത്ത ചെറിയ കുട്ടി"),
        ("retired", "Retired or not working", "വിരമിച്ചു / ജോലിയില്ല")]]
    ch += [c("care", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("home", "Home remedy", "വീട്ടുവൈദ്യം"), ("pharmacy", "Medical shop", "മെഡിക്കൽ ഷോപ്പ്"),
        ("govt", "PHC or government hospital", "പി.എച്ച്.സി / സർക്കാർ ആശുപത്രി"),
        ("private", "Private clinic or hospital", "സ്വകാര്യ ആശുപത്രി"), ("admitted", "Admitted to hospital", "ആശുപത്രിയിൽ പ്രവേശിപ്പിച്ചു")]]
    return ch


def _settings(form_id, title, instance):
    return {"form_title": title, "form_id": form_id, "version": "2026102501", "default_language": "English (en)",
            "instance_name": instance}


PERSON_ENTITY = {"list_name": "persons", "label": "concat(${new_name}, ' (', ${new_age}, ')')",
                 "create_if": "${consent_named}='yes' and ${in_list_eff}='no'"}


def form_report_diarrhoea():
    s = report_start("Report diarrhoeal disease", "വയറിളക്ക രോഗം റിപ്പോർട്ട് ചെയ്യുക") + report_person()
    s += [
        q("begin_group", "dd", "The illness", "അസുഖം", app="field-list"),
        q("integer", "stools_day", "Loose stools a day, at the worst", "ഏറ്റവും കൂടുതൽ ഉണ്ടായ ദിവസം എത്ര തവണ വയറിളക്കം?",
          req="yes", con=". >= 1 and . <= 60"),
        q("select_one yes_no", "blood", "Blood in the stool?", "മലത്തിൽ രക്തമുണ്ടോ?", req="yes", app="horizontal-compact"),
        q("select_one yes_no", "vomiting", "Vomiting?", "ഛർദ്ദിയുണ്ടോ?", req="yes", app="horizontal-compact"),
        q("select_multiple dehyd", "dehyd", "Danger signs now", "ഇപ്പോഴുള്ള അപകട ലക്ഷണങ്ങൾ", req="yes",
          con=NONE_EXCL, cmsg_en=NONE_MSG_EN, cmsg_ml=NONE_MSG_ML),
        q("end_group", "dd"),
        q("note", "refer_now", "Danger sign: give ORS now and take the person to the PHC or hospital today.",
          "അപകട ലക്ഷണം: ഇപ്പോൾ തന്നെ ഒ.ആർ.എസ് നൽകി ഇന്നുതന്നെ പി.എച്ച്.സി / ആശുപത്രിയിൽ എത്തിക്കുക.", rel=DANGER_NONE),
        q("begin_group", "dd_hh", "This household", "ഈ വീട്", app="field-list"),
        q("integer", "others_ill", "Others in this house with loose stools this week", "ഈ ആഴ്ച ഈ വീട്ടിൽ വയറിളക്കമുള്ള മറ്റുള്ളവർ",
          req="yes", default="0", con=". >= 0 and . <= 30"),
        q("select_one water_src", "water_src", "Drinking water used in the past week", "കഴിഞ്ഞ ആഴ്ച ഉപയോഗിച്ച കുടിവെള്ളം",
          req="yes"),
        q("select_one yes_no", "ors_given", "ORS given?", "ഒ.ആർ.എസ് നൽകിയോ?", req="yes", app="horizontal-compact"),
        q("end_group", "dd_hh"),
        q("calculate", "condition", calc="if(${blood}='yes', 'dysentery', 'diarrhoea')"),
        q("calculate", "danger", calc=f"if({DANGER_NONE}, 'yes', 'no')"),
    ]
    s += report_impact()
    ch = report_choices()
    ch += [c("dehyd", n, en, ml) for n, en, ml in [
        ("none", "None", "ഒന്നുമില്ല"), ("unable_drink", "Cannot drink or keeps vomiting everything", "കുടിക്കാനാകുന്നില്ല / എല്ലാം ഛർദ്ദിക്കുന്നു"),
        ("lethargic", "Very sleepy or hard to wake", "കടുത്ത മയക്കം / ഉണർത്താൻ പ്രയാസം"), ("sunken_eyes", "Sunken eyes", "കുഴിഞ്ഞ കണ്ണുകൾ"),
        ("little_urine", "Very little urine", "മൂത്രം വളരെ കുറവ്"), ("very_thirsty", "Very thirsty", "കടുത്ത ദാഹം")]]
    ch += [c("water_src", n, en, ml) for n, en, ml in [
        ("own_well", "Own well", "സ്വന്തം കിണർ"), ("shared_well", "Shared or public well", "പൊതു / പങ്കിട്ട കിണർ"),
        ("tap", "Tap (KWA / piped)", "പൈപ്പ് വെള്ളം (KWA)"), ("tanker", "Tanker", "ടാങ്കർ വെള്ളം"),
        ("borewell", "Borewell", "കുഴൽക്കിണർ"), ("packaged", "Bottled or can water", "കുപ്പി / കാൻ വെള്ളം"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    return s, ch, _settings("obs_report_diarrhoea", "Report: diarrhoeal disease",
                            "concat(${hh_id}, ' · diarrhoea · ', ${today})"), PERSON_ENTITY


HEALTH_CONDITIONS = [
    ("fever", "Fever", "പനി"), ("jaundice", "Yellow eyes (jaundice)", "മഞ്ഞപ്പിത്തം (കണ്ണിന് മഞ്ഞനിറം)"),
    ("typhoid", "Typhoid, diagnosed", "ടൈഫോയ്ഡ് (സ്ഥിരീകരിച്ചത്)"),
    ("lepto", "Leptospirosis, diagnosed or suspected", "എലിപ്പനി (സ്ഥിരീകരിച്ചതോ സംശയിക്കുന്നതോ)"),
    ("vector", "Dengue, malaria or chikungunya", "ഡെങ്കി / മലേറിയ / ചിക്കുൻഗുനിയ"),
    ("skin_water", "Skin infection after water contact", "വെള്ളം തട്ടിയ ശേഷം ത്വക്ക് അണുബാധ"),
    ("eye", "Red or sticky eyes", "കണ്ണിന് ചുവപ്പ് / പഴുപ്പ്"),
    ("diarrhoea", "Diarrhoea (the diarrhoea form asks more)", "വയറിളക്കം (വയറിളക്ക ഫോമിൽ കൂടുതൽ ചോദ്യങ്ങൾ)"),
    ("resp", "Cough, cold or breathing problem", "ചുമ / ജലദോഷം / ശ്വാസതടസ്സം"), ("heat", "Heat illness", "ചൂട് മൂലമുള്ള അസുഖം"),
    ("injury", "Injury or fall", "പരിക്ക് / വീഴ്ച"), ("chronic", "BP, sugar or other long-term illness", "ബി.പി., പ്രമേഹം, മറ്റ് ദീർഘകാല രോഗം"),
    ("mental", "Mood or memory", "മാനസികാവസ്ഥ / ഓർമ്മ"), ("other", "Other", "മറ്റുള്ളവ")]


def form_report_health():
    s = report_start("Report a health issue", "ആരോഗ്യ പ്രശ്നം റിപ്പോർട്ട് ചെയ്യുക") + report_person()
    s += [
        q("begin_group", "hi", "The health issue", "ആരോഗ്യ പ്രശ്നം", app="field-list"),
        q("select_one condition", "condition", "What is the main problem?", "പ്രധാന പ്രശ്നം എന്താണ്?", req="yes"),
        q("text", "condition_other", "Describe it", "വിവരിക്കുക", hen=VOICE_EN, hml=VOICE_ML, req="yes",
          rel="${condition}='other'"),
        q("select_one yes_no", "water_contact", "Contact with flood, canal or drain water in the 2 weeks before?",
          "അതിനു മുമ്പുള്ള 2 ആഴ്ചയിൽ വെള്ളപ്പൊക്ക / കനാൽ / ഓട വെള്ളവുമായി സമ്പർക്കം ഉണ്ടായോ?", req="yes",
          app="horizontal-compact"),
        q("end_group", "hi"),
        q("note", "refer_now", "Refer today: jaundice, suspected leptospirosis and heat illness need a doctor quickly.",
          "ഇന്നുതന്നെ റഫർ ചെയ്യുക: മഞ്ഞപ്പിത്തം, എലിപ്പനി സംശയം, ചൂട് മൂലമുള്ള അസുഖം എന്നിവയ്ക്ക് വേഗം ഡോക്ടറെ കാണണം.",
          rel="${condition}='jaundice' or ${condition}='lepto' or ${condition}='heat'"),
        q("calculate", "danger", calc="if(${condition}='jaundice' or ${condition}='lepto' or ${condition}='heat', 'yes', 'no')"),
    ]
    s += report_impact()
    ch = report_choices()
    ch += [c("condition", n, en, ml) for n, en, ml in HEALTH_CONDITIONS]
    return s, ch, _settings("obs_report_health", "Report: health issue",
                            "concat(${hh_id}, ' · ', ${condition}, ' · ', ${today})"), PERSON_ENTITY


def form_report_onehealth():
    s = [
        q("start", "start"), q("end", "end"), q("today", "today"), q("deviceid", "deviceid"),
        q("note", "intro", "Report a One Health event: something about animals, water or the surroundings that could "
          "affect health. Do not photograph people. (Malayalam text is a machine-assisted translation for "
          "accessibility; English is the reference.)",
          "വൺ ഹെൽത്ത് സംഭവം റിപ്പോർട്ട് ചെയ്യുക: ആരോഗ്യത്തെ ബാധിക്കാവുന്ന മൃഗങ്ങൾ, വെള്ളം, ചുറ്റുപാട് എന്നിവ സംബന്ധിച്ച കാര്യങ്ങൾ. "
          "ആളുകളുടെ ഫോട്ടോ എടുക്കരുത്. (മലയാളം യന്ത്രസഹായത്തോടെയുള്ള വിവർത്തനമാണ്; ഇംഗ്ലീഷ് പാഠമാണ് ആധികാരികം.)"),
        q("select_one report_role", "role", "You are reporting as", "നിങ്ങൾ റിപ്പോർട്ട് ചെയ്യുന്നത്", req="yes",
          default="${last-saved#role}"),
        q("select_one where", "where", "Where is it?", "എവിടെയാണ്?", req="yes", app="horizontal-compact"),
    ]
    s += hh_scan_block(rel="${where}='household'")
    s += [
        q("select_one_from_file sites.csv", "site_pick", "Site", "സ്ഥലം", req="yes", rel="${where}='place'",
          default="${last-saved#site_pick}"),
        q("calculate", "camp_site", calc="if(${where}='household', substring-before(${hh_id}, '-'), ${site_pick})"),
        q("begin_group", "ev", "What did you see?", "എന്താണ് കണ്ടത്?", app="field-list"),
        q("select_multiple event", "event", "What happened?", "എന്താണ് സംഭവിച്ചത്?", req="yes"),
        q("select_multiple species", "species", "Which animals or birds?", "ഏത് മൃഗങ്ങൾ / പക്ഷികൾ?", req="yes",
          rel="selected(${event}, 'animal_sick') or selected(${event}, 'animal_died')"),
        q("integer", "n_animals", "About how many?", "ഏകദേശം എത്ര?", con=". >= 1 and . <= 100000",
          rel="selected(${event}, 'animal_sick') or selected(${event}, 'animal_died') or selected(${event}, 'fish_kill')"),
        q("select_one ynd", "people_ill", "Are people nearby also ill?", "സമീപത്തുള്ള ആളുകൾക്കും അസുഖമുണ്ടോ?", req="yes",
          app="horizontal-compact"),
        q("date", "first_seen", "When was it first noticed?", "ആദ്യം ശ്രദ്ധിച്ചത് എപ്പോൾ?", req="yes", con=". <= today()"),
        q("end_group", "ev"),
        q("text", "description", "Describe what you saw", "കണ്ടത് വിവരിക്കുക", hen=VOICE_EN, hml=VOICE_ML, app="multiline"),
        q("image", "photo", "Photo (no people in it)", "ഫോട്ടോ (ആളുകൾ ഇല്ലാതെ)", params="max-pixels=1024"),
        q("geopoint", "gps", "Location", "സ്ഥാനം"),
    ]
    ch = common_choices()
    ch += [c("report_role", n, en, ml) for n, en, ml in REPORT_ROLES]
    ch += [c("where", "household", "A household", "ഒരു വീട്"), c("where", "place", "Another place", "മറ്റൊരു സ്ഥലം")]
    ch += [c("event", n, en, ml) for n, en, ml in [
        ("animal_sick", "Animals or birds sick", "മൃഗങ്ങൾ / പക്ഷികൾ രോഗബാധിതർ"),
        ("animal_died", "Animals or birds died suddenly", "മൃഗങ്ങൾ / പക്ഷികൾ പെട്ടെന്ന് ചത്തു"),
        ("fish_kill", "Dead fish in a canal, pond or the sea", "കനാലിൽ / കുളത്തിൽ / കടലിൽ ചത്ത മീനുകൾ"),
        ("rats", "Many rats", "ധാരാളം എലികൾ"), ("mosquito", "Mosquito breeding or stagnant water", "കൊതുക് വളരുന്ന കെട്ടിക്കിടക്കുന്ന വെള്ളം"),
        ("flood", "Water entered houses", "വീടുകളിൽ വെള്ളം കയറി"),
        ("water_change", "Well or tap water changed colour, smell or taste", "കിണർ / പൈപ്പ് വെള്ളത്തിന്റെ നിറം, മണം, രുചി മാറി"),
        ("sewage", "Sewage overflow or leaking septic tank", "മലിനജലം കവിഞ്ഞൊഴുകുന്നു / സെപ്റ്റിക് ടാങ്ക് ചോരുന്നു"),
        ("garbage", "Garbage dumped near water", "വെള്ളത്തിനടുത്ത് മാലിന്യം തള്ളിയിരിക്കുന്നു"), ("other", "Other", "മറ്റുള്ളവ")]]
    ch += [c("species", n, en, ml) for n, en, ml in [
        ("dog", "Dogs", "നായ്ക്കൾ"), ("cat", "Cats", "പൂച്ചകൾ"), ("cattle", "Cattle", "കന്നുകാലികൾ"), ("goat", "Goats", "ആടുകൾ"),
        ("poultry", "Hens", "കോഴികൾ"), ("duck", "Ducks", "താറാവുകൾ"), ("crow", "Crows or wild birds", "കാക്ക / മറ്റ് പക്ഷികൾ"),
        ("other", "Other", "മറ്റുള്ളവ")]]
    return s, ch, _settings("obs_report_onehealth", "Report: One Health event",
                            "concat(${camp_site}, ' · ', ${today})")


FORMS = {
    "obs_camp_registration": form_registration,
    "obs_camp_station": form_station,
    "obs_household_visit": form_home_visit,
    "obs_water_sample": form_water_sample,
    "obs_report_diarrhoea": form_report_diarrhoea,
    "obs_report_health": form_report_health,
    "obs_report_onehealth": form_report_onehealth,
}

if __name__ == "__main__":
    import os
    os.makedirs(OUT, exist_ok=True)
    for form_id, fn in FORMS.items():
        write(os.path.join(OUT, f"{form_id}.xlsx"), *fn())
    # A starting point for admins who want a new report form: a copy of the health-issue form.
    s, ch, st, ent = form_report_health()
    st = dict(st, form_id="obs_report_template", form_title="Report: [your title]")
    write(os.path.join(OUT, "report_form_template.xlsx"), s, ch, st, ent)
    print("written", len(FORMS))
